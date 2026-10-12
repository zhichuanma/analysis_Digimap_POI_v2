"""Project paths and read-only byte-level verification of source copies."""

from contextlib import contextmanager
from contextlib import ExitStack
from collections import Counter
from dataclasses import dataclass
import csv
from datetime import datetime
import hashlib
import io
from itertools import islice
import json
import math
import os
from pathlib import Path, PureWindowsPath
import re
import stat
import sqlite3
import struct
from types import MappingProxyType


def read_pua_definitions(relative_path, *, data_root, sheet=None):
    """Read an explicit CSV/XLSX table without normalizing names or evaluating formulas."""
    required = {"pua_name", "local_authority"}
    rows = []
    with open_source(relative_path, data_root=data_root) as stream:
        if Path(relative_path).suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
            try:
                if sheet not in book.sheetnames:
                    raise ValueError(f"explicit PUA sheet required; available: {book.sheetnames}")
                for cells in book[sheet].iter_rows():
                    if any(cell.data_type in {"f", "e"} for cell in cells):
                        raise ValueError("PUA definitions cannot contain formulas or error cells")
                    rows.append([cell.value for cell in cells])
            finally:
                book.close()
        elif Path(relative_path).suffix.lower() == ".csv" and sheet is None:
            text = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
            try:
                rows = list(csv.reader(text, strict=True))
            finally:
                text.detach()
        else:
            raise ValueError("PUA input requires CSV (no sheet) or XLSX (explicit sheet)")
    if not rows:
        raise ValueError("PUA table is empty")
    header = rows.pop(0)
    if (any(not isinstance(name, str) or not name for name in header)
            or len(header) != len(set(header)) or not required <= set(header)):
        raise ValueError("PUA table requires unique headers and pua_name/local_authority")
    result = []
    for number, values in enumerate(rows, 2):
        if len(values) != len(header):
            raise ValueError(f"PUA row {number}: wrong field count")
        raw = dict(zip(header, values))
        result.append({"source_path": relative_path, "source_sheet": sheet,
                       "source_row": number, "raw_fields": raw,
                       "pua_name": raw["pua_name"], "local_authority": raw["local_authority"],
                       "lad_code": raw.get("lad_code")})
    return result


def read_lad_geometries(relative_path, *, data_root, layer, code_field, name_field):
    """Read complete, valid, 2D LAD polygons in verified BNG; never repair them."""
    from shapely import from_wkb, get_coordinates
    from shapely.validation import explain_validity

    fields = _field_mapping({"code": code_field, "name": name_field}, {"code": "code", "name": "name"})
    records = []
    with _gpkg_connection(relative_path, data_root) as connection:
        # This schema helper checks metadata/CRS/keys, not Point geometry contents.
        try:
            schema = _gpkg_schema(connection, layer, fields)
        except ValueError as error:
            raise ValueError(str(error).replace("POI", "LAD")) from error
        declaration = schema["geometry"]
        declared_type = declaration["geometry_type_name"].upper()
        if declared_type not in {"POLYGON", "MULTIPOLYGON"} or declaration["z"] != 0 or declaration["m"] != 0:
            raise ValueError("LAD geometry declaration must be 2D Polygon/MultiPolygon")
        sql = f"SELECT * FROM {_quoted(layer)} ORDER BY {_quoted(schema['fid_column'])}"
        for source in connection.execute(sql):
            row = dict(source)
            fid = row[schema["fid_column"]]
            blob = row.pop(declaration["column_name"])
            try:
                if not isinstance(blob, bytes) or len(blob) < 8 or blob[:3] != b"GP\x00":
                    raise ValueError("not a version-0 GeoPackage geometry")
                flags, envelope = blob[3], (blob[3] >> 1) & 7
                if flags & 0xE0 or envelope not in (0, 1, 2, 3, 4):
                    raise ValueError("unsupported GeoPackage flags/envelope")
                srs = struct.unpack(("<" if flags & 1 else ">") + "i", blob[4:8])[0]
                if srs != declaration["srs_id"]:
                    raise ValueError("geometry header CRS differs from layer CRS")
                geometry = from_wkb(blob[8 + (0, 32, 48, 48, 64)[envelope]:])
                if bool(flags & 16) != bool(geometry.is_empty):
                    raise ValueError("GeoPackage empty flag disagrees with WKB")
                if geometry.geom_type.upper() != declared_type or geometry.has_z or geometry.has_m:
                    raise ValueError("geometry disagrees with 2D polygon declaration")
                if (geometry.is_empty or not geometry.is_valid
                        or any(not math.isfinite(v) for xy in get_coordinates(geometry) for v in xy)):
                    raise ValueError(f"empty/invalid/nonfinite LAD: {explain_validity(geometry)}")
            except Exception as error:
                raise ValueError(f"LAD FID {fid}: {error}") from error
            records.append({"source_fid": fid, "lad_code": row[code_field],
                            "lad_name": row[name_field], "raw_fields": row, "geometry": geometry})
    return records, schema


def station_rows(relative_path, *, data_root, fields=None):
    """Yield every logical CSV record with original values, including malformed rows."""
    mapping = _field_mapping(fields, STATION_FIELDS)
    with open_source(relative_path, data_root=data_root) as stream:
        text, reader, header = _station_reader(stream, mapping)
        try:
            for number, values in enumerate(reader, 1):
                raw = {name: values[i] if i < len(values) else None for i, name in enumerate(header)}
                yield {"source_record_no": number, "raw_fields": raw, "raw_values": values,
                       "raw_id": raw[mapping["id"]], "raw_longitude": raw[mapping["longitude"]],
                       "raw_latitude": raw[mapping["latitude"]],
                       "coordinate_reasons": ["csv_field_count"] if len(values) != len(header) else []}
        finally:
            text.detach()


def load_station_inspection(inspection_dir, *, data_root, station_path):
    """Reuse verified whole-source issues, including conflicts outside the target PUA."""
    root = project_path(".", root=data_root)
    directory = _audit_path(inspection_dir, root)
    evidence_path = _audit_path(str((directory.parent / "verification.json").relative_to(root)), root)
    evidence_before = _artifact_fingerprint(evidence_path)
    with _open_regular(evidence_path) as stream:
        evidence = json.load(stream, object_pairs_hook=_unique_json_object)
    if evidence.get("status") != "passed":
        raise ValueError("prior station verification is not passed")
    fingerprints = {}
    for name in ("input_inspection.json", "station_quality_issues.csv"):
        path = _audit_path(str((directory / name).relative_to(root)), root)
        actual = _artifact_fingerprint(path)
        if actual != evidence.get("artifacts", {}).get(str(path.relative_to(evidence_path.parent))):
            raise ValueError(f"prior station artifact checksum mismatch: {name}")
        fingerprints[name] = actual
    with _open_regular(directory / "input_inspection.json") as stream:
        report = json.load(stream, object_pairs_hook=_unique_json_object)
    if (report.get("status") != "completed" or report.get("scope") != "full_source"
            or any(report.get(flag) is not True for flag in
                   ("scan_complete", "full_source_complete", "duplicate_list_complete", "input_bytes_unchanged"))):
        raise ValueError("prior station audit must be completed full_source")
    info = report["stations"]
    current = _fingerprint(station_path, root, 1048576)
    if (info["path"] != station_path or info["field_mapping"] != STATION_FIELDS
            or any(current != report[key].get(station_path) for key in
                   ("input_fingerprints_before", "input_fingerprints_after"))):
        raise ValueError("current station source differs from prior audit")
    count = info["source_count"]
    if type(count) is not int or count < 0 or count != info["scanned_count"]:
        raise ValueError("prior station source counts disagree")
    issues, counts = {}, Counter()
    with _open_regular(directory / "station_quality_issues.csv") as stream:
        text = io.TextIOWrapper(stream, encoding="utf-8", newline="")
        try:
            reader = csv.DictReader(text, strict=True)
            if reader.fieldnames != ["source_record_no", "StationID", "reason", "detail"]:
                raise ValueError("invalid prior station issue schema")
            for row in reader:
                number = int(row["source_record_no"])
                if not 1 <= number <= count or not row["reason"] or None in row:
                    raise ValueError("invalid prior station issue row")
                issues.setdefault(number, []).append(row)
                counts[row["reason"]] += 1
        finally:
            text.detach()
    if dict(counts) != info["issue_counts"]:
        raise ValueError("prior station issue totals disagree")
    if (any(value != _artifact_fingerprint(directory / name) for name, value in fingerprints.items())
            or evidence_before != _artifact_fingerprint(evidence_path)):
        raise ValueError("prior station evidence changed while reading")
    return {"source_count": count, "issues": issues, "groups": info["groups"],
            "inspection_dir": inspection_dir, "artifacts": fingerprints,
            "verification": evidence_before, "station_fingerprint": current}


def _read_json(path):
    with _open_regular(path) as stream:
        return json.load(stream, object_pairs_hook=_unique_json_object, parse_constant=_reject_json_constant)


def load_station_scope(scope_dir, *, data_root, station_path, coordinate_operation):
    """Read sealed, successful step-09 output and reconcile every raw station row."""
    from shapely import from_wkb
    from src.geography import _polygon, _project_rows
    from src.spatial import validate_coordinate_operation

    root = project_path(".", root=data_root)
    directory = _audit_path(scope_dir, root)
    evidence_path = _audit_path(str((directory.parent / "verification.json").relative_to(root)), root)
    evidence_fp = _artifact_fingerprint(evidence_path)
    evidence = _read_json(evidence_path)
    if evidence.get("status") != "passed":
        raise ValueError("step09 verification must be passed")
    names = ("pua_scope_report.json", "scope_definition.json", "study_area.wkb",
             "station_pua_audit.csv", "stations_in_target_pua.csv")
    fingerprints = {}
    for name in names:
        path = _audit_path(str((directory / name).relative_to(root)), root)
        actual = _artifact_fingerprint(path)
        if actual != evidence.get("artifacts", {}).get(str(path.relative_to(evidence_path.parent))):
            raise ValueError(f"step09 artifact checksum mismatch: {name}")
        fingerprints[name] = actual
    report = _read_json(directory / names[0])
    if report.get("status") != "completed" or any(report.get(k) is not True for k in
            ("mapping_complete", "scope_complete", "scan_complete", "validation_complete", "input_bytes_unchanged")):
        raise ValueError("step09 scope must be completed and validated")
    for name in names[1:]:
        if fingerprints[name] != report.get("artifacts", {}).get(name):
            raise ValueError(f"step09 report artifact mismatch: {name}")
    definition = _read_json(directory / "scope_definition.json")
    if (definition.get("predicate") != "covers" or definition.get("crs") != "EPSG:27700"
            or definition.get("coordinate_operation") != validate_coordinate_operation(coordinate_operation)
            or definition.get("target_puas") != report.get("target_puas")):
        raise ValueError("step09 scope definition/CRS/pipeline/targets mismatch")
    input_fp = definition["input_fingerprints"]
    if station_path not in input_fp or input_fp != report.get("input_fingerprints_before") or input_fp != report.get("input_fingerprints_after"):
        raise ValueError("step09 source fingerprints disagree")
    for path, fp in input_fp.items():
        if _fingerprint(path, root, 1048576) != fp:
            raise ValueError(f"step09 source has changed: {path}")
    with _open_regular(directory / "study_area.wkb") as stream:
        study = _polygon(from_wkb(stream.read()), "step09 study area")
    tables = []
    for name in ("station_pua_audit.csv", "stations_in_target_pua.csv"):
        with _open_regular(directory / name) as stream:
            text = io.TextIOWrapper(stream, encoding="utf-8", newline="")
            try:
                reader = csv.DictReader(text, strict=True)
                expected = ["source_record_no", "raw_id", "raw_longitude", "raw_latitude", "raw_fields", "raw_values",
                            "xy", "lad_codes", "pua_names", "target_pua_names", "in_target_pua", "audit_reasons",
                            "prior_station_issues", "transformation_error"]
                if reader.fieldnames != expected:
                    raise ValueError("invalid step09 station audit schema")
                table = list(reader)
                if any(None in r or any(v is None for v in r.values()) for r in table):
                    raise ValueError("malformed step09 station audit row")
                tables.append(table)
            finally:
                text.detach()
    audit, candidates = tables
    if (len(audit) != report.get("source_record_count") or len(candidates) != report.get("candidate_count")
            or candidates != [r for r in audit if r["in_target_pua"] == "True"]):
        raise ValueError("step09 source/candidate rowsets disagree")
    projected = list(_project_rows(station_rows(station_path, data_root=root), coordinate_operation, definition["station_batch_size"]))
    if len(projected) != len(audit):
        raise ValueError("step09 audit is not the entire current station source")
    rows = []
    for source, saved in zip(projected, audit, strict=True):
        if saved["source_record_no"] != str(source["source_record_no"]):
            raise ValueError("step09 audit logical row sequence disagrees")
        row = {"source_record_no": source["source_record_no"]}
        for name in ("raw_fields", "raw_values", "lad_codes", "pua_names", "target_pua_names", "audit_reasons", "prior_station_issues"):
            row[name] = json.loads(saved[name], object_pairs_hook=_unique_json_object, parse_constant=_reject_json_constant)
        for name in ("raw_fields", "raw_values"):
            if row[name] != source[name]:
                raise ValueError(f"step09 raw values mismatch: {name}")
        for name in ("raw_id", "raw_longitude", "raw_latitude"):
            if saved[name] != ("" if source[name] is None else source[name]):
                raise ValueError(f"step09 raw column mismatch: {name}")
            row[name] = source[name]
        row["xy"] = json.loads(saved["xy"]) if saved["xy"] else None
        xy = list(source["xy"]) if source["xy"] is not None else None
        if row["xy"] != xy:
            raise ValueError("step09 projected coordinates differ from fixed pipeline")
        if saved["in_target_pua"] not in ("True", "False", ""):
            raise ValueError("invalid step09 geography tri-state")
        row["in_target_pua"] = {"True": True, "False": False, "": None}[saved["in_target_pua"]]
        if xy is not None:
            from shapely.geometry import Point
            if study.covers(Point(xy)) != row["in_target_pua"]:
                raise ValueError("step09 position/study membership mismatch")
        elif row["in_target_pua"] is not None:
            raise ValueError("unlocated step09 row has known membership")
        row["coordinate_reasons"] = source["coordinate_reasons"]
        row["transformation_error"] = saved["transformation_error"] or None
        rows.append(row)
    if (any(fp != _artifact_fingerprint(directory / name) for name, fp in fingerprints.items())
            or evidence_fp != _artifact_fingerprint(evidence_path)):
        raise ValueError("step09 evidence changed while reading")
    return {"rows": rows, "study_area": study, "definition": definition,
            "provenance": {"scope_dir": scope_dir, "artifacts": fingerprints, "verification": evidence_fp,
                           "input_fingerprints": input_fp, "candidate_count": len(candidates)}}


def read_order_extent(order_path, poi_path, *, data_root, poi_layer, review):
    """Read one identified delivery rectangle; never infer extent from POI points.

    The catalogue lacks a CRS label. A required, saved review connects product
    National Grid documentation to the identified GPKG's verified metre CRS.
    Multiple folders/chunks, malformed or disagreeing extents fail explicitly.
    """
    if (not isinstance(review, dict) or review.get("crs") != "EPSG:27700" or review.get("units") != "metre"
            or review.get("axis_order") != ["easting", "northing"] or not review.get("references")):
        raise ValueError("explicit order CRS/units/axis review required")
    for reference in review["references"]:
        if not all(isinstance(reference.get(k), str) and reference[k].strip() for k in ("source_path", "locator", "quote")):
            raise ValueError("order review requires source path, locator and quote")
    with open_source(order_path, data_root=data_root) as stream:
        text = stream.read().decode("utf-8-sig")
    entries = {}
    for line_number, line in enumerate(text.splitlines(), 1):
        match = re.fullmatch(r"\s*(Folder|Filename|Extent|Product|Version|Format|Date Ordered|Area):\s*(.*?)\s*", line)
        if match:
            entries.setdefault(match[1], []).append({"line": line_number, "text": line, "value": match[2]})
    filename = Path(poi_path).name
    folder = Path(filename).stem
    required = {"Folder": folder, "Filename": f"{folder}/{filename}", "Product": "Points of Interest", "Format": "GeoPackage"}
    for field, expected in required.items():
        if len(entries.get(field, [])) != 1 or entries[field][0]["value"] != expected:
            raise ValueError(f"order delivery identity is ambiguous or mismatched: {field}")
    if any(len(entries.get(field, [])) != 1 for field in ("Version", "Date Ordered", "Area")) or len(entries.get("Extent", [])) != 2:
        raise ValueError("order requires one folder and one chunk with two matching extents")
    numeric = r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?"
    pattern = rf"({numeric})\s*,\s*({numeric})\s*-\s*({numeric})\s*,\s*({numeric})"
    rectangles = []
    for extent in entries["Extent"]:
        match = re.fullmatch(pattern, extent["value"])
        if match is None:
            raise ValueError("malformed order Extent")
        bounds = [float(v) for v in match.groups()]
        if any(not math.isfinite(v) for v in bounds) or bounds[0] >= bounds[2] or bounds[1] >= bounds[3]:
            raise ValueError("invalid order extent values")
        rectangles.append(bounds)
    if rectangles[0] != rectangles[1]:
        raise ValueError("folder/chunk order extents disagree")
    bounds = rectangles[0]
    area = re.fullmatch(rf"({numeric}) square kilometres", entries["Area"][0]["value"])
    if area is None or float(area[1]) != (bounds[2] - bounds[0]) * (bounds[3] - bounds[1]) / 1e6:
        raise ValueError("order area does not agree with metre rectangle")
    with _gpkg_connection(poi_path, data_root) as connection:
        metadata = _gpkg_schema(connection, poi_layer, POI_FIELDS)
    return {"bounds": bounds, "crs": "EPSG:27700", "units": "metre", "axis_order": ["easting", "northing"],
            "source_path": order_path, "poi_path": poi_path, "poi_layer": poi_layer,
            "catalogue_entries": entries, "review": review, "poi_metadata": metadata,
            "limitation": "catalogue has no CRS label; National Grid product review plus identified delivery CRS; coverage is a proxy"}


def station_parquet_schemas():
    """Explicit schema also applies to empty/all-null station preparations."""
    import pyarrow as pa
    text, number, flag = pa.string(), pa.float64(), pa.bool_()
    # Parquet's compliant list encoding names its child "element". Declare it
    # explicitly so metadata-aware equality also holds after file round-trip.
    strings, integers = pa.list_(pa.field("element", text)), pa.list_(pa.field("element", pa.int64()))
    station = pa.schema([
        ("station_key", text), ("source_record_nos", integers), ("raw_ids", strings), ("normalized_ids", strings),
        ("source_row_count", pa.int64()), ("entity_resolved", flag), ("longitude", number), ("latitude", number),
        ("xy", pa.list_(pa.field("element", number))), ("lad_codes", strings), ("pua_names", strings), ("target_pua_names", strings),
        ("in_target_pua", flag), ("selected", flag), ("identity_eligible", flag), ("identity_reasons", strings),
        ("audit_reasons", strings), ("within_gb_proxy", flag), ("within_ni_proxy", flag), ("within_order", flag),
        ("gb_boundary_distance_m", number), ("order_clearance_m", number), ("spatial_status", text),
        ("query_eligible", flag), ("main_status", text), ("query_blockers", strings),
    ], metadata={b"schema_version": b"1", b"crs": b"EPSG:27700", b"step": b"10"})
    radius = pa.schema([
        ("station_key", text), ("radius_m", pa.int64()), ("selected", flag), ("query_eligible", flag),
        ("main_status", text), ("within_gb_proxy", flag), ("gb_boundary_distance_m", number),
        ("order_clearance_m", number), ("coverage_complete_proxy", flag), ("coverage_reasons", strings),
    ], metadata={b"schema_version": b"1", b"distance_units": b"metre", b"step": b"10"})
    return station, radius


def _write_parquet(path, rows, schema):
    import pyarrow as pa
    import pyarrow.parquet as pq
    if any(set(row) != set(schema.names) for row in rows):
        raise ValueError("Parquet record fields disagree with explicit schema")
    table = pa.Table.from_pylist(rows, schema=schema)
    with path.open("xb") as stream:
        pq.write_table(table, stream)
    restored = pq.read_table(path)
    if not restored.schema.equals(schema, check_metadata=True) or restored.to_pylist() != rows:
        raise ValueError(f"Parquet read-back differs: {path.name}")


class StationPreparationError(ValueError):
    """Partial files beside a failed report are diagnostic, not valid outputs."""

    def __init__(self, report):
        self.report = report
        super().__init__("; ".join(report["errors"]))


def prepare_station_files(station_path, lad_path, order_path, poi_path, *, data_root, output_dir,
                          scope_dir, inspection_dir, coordinate_operation, required_paths,
                          order_review, poi_layer, sample_size=None, random_seed=None):
    """Run only step 10, always checking full source before any entity sampling."""
    from shapely import from_wkb
    from src.geography import _cell, _write_table, coverage_geometries
    from src.spatial import prepare_stations, station_radius_base, validate_coordinate_operation

    root = project_path(".", root=data_root)
    output = project_path(output_dir, root=root)
    if (not Path(output_dir).is_relative_to("outputs") or output == root / "outputs"
            or not output.is_relative_to(root / "outputs")):
        raise ValueError("output_dir must be a new directory strictly inside outputs/")
    output.mkdir(parents=True, exist_ok=False)
    report = {"schema_version": 1, "status": "running", "errors": [], "output_dir": str(output),
              "scope": "sample" if sample_size is not None else "full_source", "sample_size": sample_size,
              "random_seed": random_seed, "validation_complete": False, "poi_matching_performed": False}
    _json_write(output / "station_preparation_report.json", report)
    try:
        operation = validate_coordinate_operation(coordinate_operation)
        paths = list(required_paths)
        references = {r["source_path"] for r in order_review.get("references", [])}
        if not {station_path, lad_path, order_path, poi_path} | references <= set(paths):
            raise ValueError("required_paths must include all consumed sources and order references")
        report["sources_before"] = verify_sources(data_root=root, required_paths=paths)
        before = {p: _fingerprint(p, root, 1048576) for p in paths}
        report["input_fingerprints_before"] = before
        upstream = load_station_scope(scope_dir, data_root=root, station_path=station_path, coordinate_operation=operation)
        definition = upstream["definition"]
        if not set(definition["required_paths"]) <= set(paths) or definition["lad_path"] != lad_path:
            raise ValueError("required_paths/LAD must include the verified step09 sources")
        prior = load_station_inspection(inspection_dir, data_root=root, station_path=station_path)
        if prior["source_count"] != len(upstream["rows"]):
            raise ValueError("full source count differs from step06")
        for row in upstream["rows"]:
            if row["prior_station_issues"] != prior["issues"].get(row["source_record_no"], []):
                raise ValueError("step09 row issues differ from full-source step06 evidence")
        report["upstream_scope"] = upstream["provenance"]
        report["prior_station_audit"] = {k: v for k, v in prior.items() if k != "issues"}
        lads, lad_metadata = read_lad_geometries(lad_path, data_root=root, layer=definition["lad_layer"],
                                                code_field=definition["lad_code_field"], name_field=definition["lad_name_field"])
        order = read_order_extent(order_path, poi_path, data_root=root, poi_layer=poi_layer, review=order_review)
        coverage = coverage_geometries(lads, crs=lad_metadata["crs"]["definition"], order_bounds=order["bounds"])
        if not coverage["gb"].covers(upstream["study_area"]):
            raise ValueError("step09 study area is not contained in full GB proxy")
        prepared = prepare_stations(upstream["rows"], coverage, target_puas=definition["target_puas"],
                                    sample_size=sample_size, random_seed=random_seed)
        stations, audit = prepared["stations"], prepared["audit"]
        base = station_radius_base(stations)
        station_schema, radius_schema = station_parquet_schemas()
        _write_parquet(output / "stations.parquet", stations, station_schema)
        _write_parquet(output / "station_radius_base.parquet", base, radius_schema)
        fields = ["source_record_no", "raw_id", "raw_longitude", "raw_latitude", "raw_fields", "raw_values", "xy",
                  "lad_codes", "pua_names", "target_pua_names", "in_target_pua", "coordinate_reasons", "audit_reasons",
                  "prior_station_issues", "transformation_error", "normalized_id", "identity_reasons", "station_key",
                  "id_source_record_nos", "merge_action", "selected", "query_eligible", "main_status",
                  "entity_query_blockers", "entity_audit_reasons"]
        _write_table(output / "station_audit.csv", audit, fields)
        with (output / "station_audit.csv").open(encoding="utf-8", newline="") as stream:
            restored = list(csv.DictReader(stream, strict=True))
        expected = [{k: "" if r.get(k) is None else str(_cell(r[k])) for k in fields} for r in audit]
        if restored != expected:
            raise ValueError("station audit CSV read-back differs")
        for key, name in (("gb", "gb_proxy.wkb"), ("ni", "ni_proxy.wkb"), ("data_coverage", "data_coverage_proxy.wkb")):
            with (output / name).open("xb") as stream:
                stream.write(coverage[key].wkb)
            if not from_wkb((output / name).read_bytes()).equals(coverage[key]):
                raise ValueError("coverage WKB read-back differs")
        coverage_definition = {"schema_version": 1, "crs": coverage["crs"], "predicate": "covers", "units": "metre",
                               "gb_lad_codes": coverage["members"]["gb"], "ni_lad_codes": coverage["members"]["ni"],
                               "lad_metadata": lad_metadata, "order": order, "coordinate_operation": operation,
                               "target_puas": definition["target_puas"], "upstream_scope": upstream["provenance"],
                               "input_fingerprints": before, "station_scope": "target PUA covers EVCS only",
                               "poi_scope": "GB/order coverage covers POI and distance <= radius; no PUA clipping; cleaning deferred",
                               "coverage_rule": "query_eligible and within_gb_proxy and gb_distance >= r and order_clearance >= r",
                               "identity_rule": "strip IDs preserving leading zeros; exact parsed WGS84 co-location; entire location blocked by missing/duplicate/conflicting ID",
                               "sample_keys": prepared["sample_keys"], "radii_m": [500, 1000],
                               "limitations": ["GB LAD proxy: generalised boundaries, holes retained; not proof of POI completeness",
                                               "fixed approximate seven-parameter station conversion; no radius tolerance",
                                               order["limitation"]]}
        _json_write(output / "coverage_definition.json", coverage_definition)
        if _read_json(output / "coverage_definition.json") != coverage_definition:
            raise ValueError("coverage definition read-back differs")
        source_keys = sorted(n for s in stations for n in s["source_record_nos"])
        if source_keys != list(range(1, len(audit) + 1)) or len(base) != 2 * sum(s["selected"] for s in stations):
            raise ValueError("source/entity/radius conservation failed")
        if len({(r["station_key"], r["radius_m"]) for r in base}) != len(base):
            raise ValueError("duplicate station/radius pair")
        report["sources_after"] = verify_sources(data_root=root, required_paths=paths)
        after = {p: _fingerprint(p, root, 1048576) for p in paths}
        check_scope = load_station_scope(scope_dir, data_root=root, station_path=station_path, coordinate_operation=operation)
        if (after != before or check_scope["provenance"] != upstream["provenance"]
                or load_station_inspection(inspection_dir, data_root=root, station_path=station_path) != prior):
            raise ValueError("input or upstream evidence changed during station preparation")
        selected = [s for s in stations if s["selected"]]
        report.update(status="completed", validation_complete=True, input_bytes_unchanged=True,
                      full_source_identity_checked=True, input_fingerprints_after=after,
                      source_record_count=len(audit), entity_record_count=len(stations),
                      resolved_entity_count=sum(s["entity_resolved"] for s in stations),
                      unresolved_record_count=sum(not s["entity_resolved"] for s in stations),
                      merged_location_count=sum(s["source_row_count"] > 1 for s in stations),
                      merged_excess_rows=len(audit) - len(stations),
                      target_source_rows=sum(r["in_target_pua"] is True for r in audit),
                      target_entity_count=sum(s["in_target_pua"] is True for s in stations),
                      selected_entity_count=len(selected), query_eligible_count=sum(s["query_eligible"] for s in selected),
                      nonqueryable_selected_count=sum(not s["query_eligible"] for s in selected), radius_row_count=len(base),
                      main_status_counts=dict(Counter(s["main_status"] for s in stations)),
                      spatial_status_counts=dict(Counter(s["spatial_status"] for s in stations)),
                      identity_reason_counts=dict(Counter(reason for r in audit for reason in r["identity_reasons"])),
                      selected_query_blocker_counts=dict(Counter(reason for s in selected for reason in s["query_blockers"])),
                      target_puas=definition["target_puas"],
                      pua_entity_counts={p: sum(s["in_target_pua"] is True and p in s["target_pua_names"] for s in stations) for p in definition["target_puas"]},
                      coverage_counts={str(r): {"complete": sum(b["coverage_complete_proxy"] for b in base if b["radius_m"] == r),
                                                 "incomplete": sum(not b["coverage_complete_proxy"] for b in base if b["radius_m"] == r)} for r in (500, 1000)},
                      gb_lad_count=len(coverage["members"]["gb"]), ni_lad_count=len(coverage["members"]["ni"]))
        report["artifacts"] = {p.name: _artifact_fingerprint(p) for p in sorted(output.iterdir()) if p.name != "station_preparation_report.json"}
    except Exception as error:
        report.update(status="failed", validation_complete=False)
        report["errors"].append(f"{type(error).__name__}: {error}")
        _json_write(output / "station_preparation_report.json", report)
        raise StationPreparationError(report) from error
    _json_write(output / "station_preparation_report.json", report)
    return report


def _resolve(path: Path) -> Path:
    """Resolve existing links, allowing new targets but rejecting link loops."""
    try:
        return path.resolve(strict=True)
    except FileNotFoundError:
        return path.resolve(strict=False)


def project_path(relative_path: str, *, root: Path, field: str = "path") -> Path:
    """Resolve a POSIX relative path inside an explicitly trusted absolute root.

    Check the original text before normalization. Neither the current working
    directory nor expansion of '~' or environment variables participates.
    Missing targets are allowed; existing parent symlinks still constrain them.
    No directories or files are created.
    """
    if not isinstance(relative_path, str) or not relative_path.strip():
        raise ValueError(f"{field}: must be a non-empty relative path string")
    if "\x00" in relative_path or "\\" in relative_path:
        raise ValueError(f"{field}: must use POSIX path syntax without NUL bytes")
    path = Path(relative_path)
    if path.is_absolute() or PureWindowsPath(relative_path).drive:
        raise ValueError(f"{field}: absolute or drive-qualified paths are forbidden")
    if ".." in path.parts:
        raise ValueError(f"{field}: '..' path components are forbidden")

    root = Path(root)
    if not root.is_absolute():
        raise ValueError(f"{field}: root must be an absolute trusted path")
    try:
        resolved_root = _resolve(root)
        resolved = _resolve(resolved_root / path)
    except (OSError, RuntimeError) as exc:
        raise ValueError(f"{field}: cannot resolve path: {exc}") from exc
    if not resolved.is_relative_to(resolved_root):
        raise ValueError(f"{field}: resolved path escapes root {resolved_root}")
    return resolved


class SourceError(ValueError):
    """An input could not be read under the source-file contract."""

    def __init__(self, code, message):
        self.code = code
        super().__init__(message)


class SourceVerificationError(ValueError):
    """A failed verification; ``report`` is JSON-serializable evidence."""

    def __init__(self, report):
        self.report = report
        errors = report["errors"] + [
            error for row in report["files"] for error in row["errors"]
        ]
        super().__init__("Source verification failed: " + "; ".join(
            error["message"] for error in errors
        ))


def _source_path(relative_path, data_root):
    try:
        root = project_path(".", root=data_root, field="data_root")
        resolved = project_path(relative_path, root=root, field="source_path")
    except (TypeError, ValueError) as exc:
        raise SourceError("invalid_source_path", str(exc)) from exc
    areas = (root / "data/raw", root / "cities")
    lexical = root / relative_path
    if not all(any(path.is_relative_to(area) for area in areas) for path in (lexical, resolved)):
        raise SourceError(
            "invalid_source_path", f"{relative_path}: inputs must stay in data/raw/ or cities/"
        )
    return resolved


@contextmanager
def _open_regular(path):
    # O_NONBLOCK prevents FIFO input from hanging before the regular-file check.
    # The final component was resolved already; do not follow a replacement link.
    try:
        descriptor = os.open(path, os.O_RDONLY | os.O_NONBLOCK | os.O_NOFOLLOW)
    except FileNotFoundError as exc:
        raise SourceError("missing_file", f"{path}: file is missing") from exc
    except OSError as exc:
        raise SourceError("unreadable_file", f"{path}: cannot open read-only: {exc}") from exc
    try:
        if not stat.S_ISREG(os.fstat(descriptor).st_mode):
            raise SourceError("not_regular_file", f"{path}: must be a regular file")
        stream = os.fdopen(descriptor, "rb")
        descriptor = None
        with stream:
            yield stream
    finally:
        if descriptor is not None:
            os.close(descriptor)


@contextmanager
def open_source(relative_path, *, data_root):
    """Yield a read-only binary stream from data/raw/ or cities/ in data_root.

    No write mode, permission change, directory creation or provenance-source
    access is provided. A source's content is not verified merely by opening it.
    """
    path = _source_path(relative_path, data_root)
    with _open_regular(path) as stream:
        yield stream


def _file_state(info):
    return info.st_dev, info.st_ino, info.st_size, info.st_mtime_ns, info.st_ctime_ns


def _fingerprint(relative_path, data_root, chunk_size):
    if type(chunk_size) is not int or chunk_size <= 0:
        raise ValueError("chunk_size: must be a positive integer, not a bool")
    path = _source_path(relative_path, data_root)
    digest = hashlib.sha256()
    size = 0
    try:
        with open_source(relative_path, data_root=data_root) as stream:
            before = os.fstat(stream.fileno())
            while block := stream.read(chunk_size):
                size += len(block)
                digest.update(block)
            after = os.fstat(stream.fileno())
            if _file_state(before) != _file_state(after) or _file_state(after) != _file_state(path.stat()):
                raise SourceError("changed_during_read", f"{relative_path}: changed during reading")
            if size != after.st_size:
                raise SourceError("changed_during_read", f"{relative_path}: inconsistent byte count")
    except OSError as exc:
        raise SourceError("unreadable_file", f"{relative_path}: cannot finish reading: {exc}") from exc
    return {"resolved_path": str(path), "size_bytes": size, "sha256": digest.hexdigest()}


def sha256(relative_path, *, data_root, chunk_size=1048576):
    """Hash all original bytes in bounded chunks; return a lowercase digest."""
    return _fingerprint(relative_path, data_root, chunk_size)["sha256"]


def _unique_json_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"duplicate JSON key: {key}")
        result[key] = value
    return result


def _reject_json_constant(value):
    raise ValueError(f"non-standard JSON constant: {value}")


def _required_paths(values, data_root):
    if not isinstance(values, (list, tuple)) or not values:
        raise ValueError("required_paths: must be an independent non-empty list or tuple")
    normalized = []
    resolved = set()
    for value in values:
        path = _source_path(value, data_root)
        name = Path(value).as_posix()
        if name in normalized or path in resolved:
            raise ValueError(f"required_paths: duplicate path or alias: {value}")
        normalized.append(name)
        resolved.add(path)
    return normalized


def _manifest_entries(document, data_root):
    if not isinstance(document, dict) or set(document) != {"schema_version", "files"}:
        raise ValueError("manifest: expected schema_version and files")
    if type(document["schema_version"]) is not int or document["schema_version"] != 1:
        raise ValueError("manifest.schema_version: must be integer 1")
    if not isinstance(document["files"], list):
        raise ValueError("manifest.files: must be a list")
    fields = {
        "path", "size_bytes", "sha256", "source_location", "copied_at",
        "source_sha256_before_copy", "source_sha256_after_copy",
    }
    entries = []
    names, resolved = set(), set()
    for index, entry in enumerate(document["files"]):
        label = f"manifest.files[{index}]"
        if not isinstance(entry, dict) or not fields.issubset(entry):
            raise ValueError(f"{label}: missing required source/copy evidence fields")
        path = _source_path(entry["path"], data_root)
        name = Path(entry["path"]).as_posix()
        if name in names or path in resolved:
            raise ValueError(f"{label}.path: duplicate path or alias: {entry['path']}")
        names.add(name)
        resolved.add(path)
        if type(entry["size_bytes"]) is not int or entry["size_bytes"] < 0:
            raise ValueError(f"{label}.size_bytes: must be a non-negative integer, not a bool")
        for field in ("sha256", "source_sha256_before_copy", "source_sha256_after_copy"):
            if not isinstance(entry[field], str) or not re.fullmatch(r"[0-9a-fA-F]{64}", entry[field]):
                raise ValueError(f"{label}.{field}: must contain exactly 64 hexadecimal digits")
        if not isinstance(entry["source_location"], str) or not entry["source_location"].strip():
            raise ValueError(f"{label}.source_location: must be non-empty provenance text")
        try:
            copied_at = datetime.fromisoformat(entry["copied_at"])
            if copied_at.tzinfo is None:
                raise ValueError("missing timezone")
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{label}.copied_at: must be an ISO 8601 timestamp with timezone") from exc
        entries.append({**entry, "path": name, "resolved_path": str(path)})
    return entries


def _issue(code, message, path=None):
    return {"code": code, "path": path, "message": message}


def verify_sources(manifest_path="data/source_manifest.json", *, data_root, required_paths):
    """Verify all declared copies against a separately supplied required list.

    Return a JSON-serializable report, or raise SourceVerificationError carrying
    the failed report. Only a new-project manifest and new-project copies are
    read. Historical source hashes are compared, never recreated or repaired.
    The caller, not this function, may persist the report in outputs/<run-id>/.
    """
    report = {
        "status": "failed", "manifest_path": None, "data_root": None,
        "required_paths": [], "declared_count": 0, "checked_count": 0,
        "passed_count": 0, "files": [], "errors": [],
    }
    try:
        root = project_path(".", root=data_root, field="data_root")
        report["data_root"] = str(root)
        path = project_path(manifest_path, root=root, field="manifest_path")
        if Path(manifest_path) != Path("data/source_manifest.json") or path != root / "data/source_manifest.json":
            raise ValueError("manifest_path: must stay at data/source_manifest.json")
        report["manifest_path"] = str(path)
    except (TypeError, ValueError) as exc:
        report["errors"].append(_issue("invalid_manifest_path", str(exc)))
        raise SourceVerificationError(report) from exc

    # Check these independent prerequisites separately so a missing required
    # list does not conceal a missing manifest, or vice versa.
    entries = []
    try:
        with _open_regular(path) as stream:
            document = json.load(
                stream, object_pairs_hook=_unique_json_object, parse_constant=_reject_json_constant
            )
        if isinstance(document, dict) and isinstance(document.get("files"), list):
            report["declared_count"] = len(document["files"])
        entries = _manifest_entries(document, root)
    except SourceError as exc:
        code = "missing_manifest" if exc.code == "missing_file" else exc.code
        report["errors"].append(_issue(code, str(exc), "data/source_manifest.json"))
    except (OSError, ValueError) as exc:
        report["errors"].append(_issue("invalid_manifest", str(exc), "data/source_manifest.json"))
    try:
        report["required_paths"] = _required_paths(required_paths, root)
    except (TypeError, ValueError) as exc:
        report["errors"].append(_issue("invalid_required_paths", str(exc)))
    if report["errors"]:
        raise SourceVerificationError(report)

    required = set(report["required_paths"])
    declared = {entry["path"] for entry in entries}
    for missing in sorted(required - declared):
        report["errors"].append(_issue("missing_declaration", f"{missing}: required but not declared", missing))

    for entry in entries:
        name = entry["path"]
        row = {
            "path": name, "resolved_path": entry["resolved_path"], "required": name in required,
            "expected_size_bytes": entry["size_bytes"], "actual_size_bytes": None,
            "expected_sha256": entry["sha256"].lower(), "actual_sha256": None,
            "source_location": entry["source_location"], "copied_at": entry["copied_at"],
            "source_sha256_before_copy": entry["source_sha256_before_copy"].lower(),
            "source_sha256_after_copy": entry["source_sha256_after_copy"].lower(),
            "status": "failed", "errors": [],
        }
        for field in ("source_sha256_before_copy", "source_sha256_after_copy"):
            if row[field] != row["expected_sha256"]:
                row["errors"].append(_issue(
                    "copy_evidence_mismatch", f"{name}: {field} differs from declared copy sha256", name
                ))
        try:
            actual = _fingerprint(name, root, 1048576)
            report["checked_count"] += 1
            row.update(resolved_path=actual["resolved_path"], actual_size_bytes=actual["size_bytes"], actual_sha256=actual["sha256"])
            if row["actual_size_bytes"] != row["expected_size_bytes"]:
                row["errors"].append(_issue("size_mismatch", f"{name}: byte count differs from manifest", name))
            if row["actual_sha256"] != row["expected_sha256"]:
                row["errors"].append(_issue("sha256_mismatch", f"{name}: SHA-256 differs from manifest", name))
        except SourceError as exc:
            row["errors"].append(_issue(exc.code, str(exc), name))
        if not row["errors"]:
            row["status"] = "passed"
            report["passed_count"] += 1
        report["files"].append(row)

    if report["errors"] or report["passed_count"] != report["declared_count"]:
        raise SourceVerificationError(report)
    report["status"] = "passed"
    return report


# Step 06 audits source records. None of these functions selects or cleans POIs.
POI_FIELDS = {
    "ref_no": "ref_no", "raw_code": "pointx_class",
    "easting": "feature_easting", "northing": "feature_northing",
    "accuracy": "pos_accuracy",
}
STATION_FIELDS = {"id": "StationID", "latitude": "Latitude", "longitude": "Longitude"}


class InputInspectionError(ValueError):
    """A failed/incomplete inspection; partial artifacts are never complete lists."""

    def __init__(self, report):
        self.report = report
        super().__init__("; ".join(report["errors"]))


def _positive_integer(value, name):
    if type(value) is not int or value <= 0:
        raise ValueError(f"{name}: must be a positive integer, not a bool")


def _quoted(name):
    return '"' + name.replace('"', '""') + '"'


def _field_mapping(value, default):
    fields = dict(default) if value is None else value
    if not isinstance(fields, dict) or fields.keys() != default.keys():
        raise ValueError(f"field mapping must specify {sorted(default)}")
    if any(not isinstance(v, str) or not v or '\x00' in v for v in fields.values()):
        raise ValueError("field mapping values must be non-empty column names")
    if len(set(fields.values())) != len(fields):
        raise ValueError("field mapping must use distinct columns")
    return dict(fields)


@contextmanager
def _gpkg_connection(relative_path, data_root):
    path = _source_path(relative_path, data_root)
    # A source snapshot must not depend on unmanifested SQLite sidecar files.
    if any(Path(str(path) + suffix).exists() for suffix in ("-wal", "-journal", "-shm")):
        raise ValueError("GPKG has SQLite sidecars; a standalone stable snapshot is required")
    with _open_regular(path) as stream:
        before = _file_state(os.fstat(stream.fileno()))
        connection = sqlite3.connect(path.as_uri() + "?mode=ro", uri=True)
        try:
            connection.row_factory = sqlite3.Row
            connection.execute("PRAGMA query_only=ON")
            connection.execute("PRAGMA trusted_schema=OFF")
            connection.execute("BEGIN")
            yield connection
        finally:
            connection.close()
            if before != _file_state(path.stat()) or before != _file_state(os.fstat(stream.fileno())):
                raise ValueError("GPKG changed during reading")


def _gpkg_schema(connection, layer, fields):
    from pyproj import CRS

    contents = [dict(row) for row in connection.execute("SELECT * FROM gpkg_contents")]
    layers = [row["table_name"] for row in contents if row["data_type"] == "features"]
    if layer is None:
        if len(layers) != 1:
            raise ValueError(f"POI layer is ambiguous; available feature layers: {layers}")
        layer = layers[0]
    if not isinstance(layer, str) or layer not in layers:
        raise ValueError(f"POI layer {layer!r} is not a registered feature layer: {layers}")
    columns = [dict(row) for row in connection.execute(f"PRAGMA table_info({_quoted(layer)})")]
    names = {column["name"] for column in columns}
    missing = set(fields.values()) - names
    if missing:
        raise ValueError(f"POI missing required fields: {sorted(missing)}")
    keys = [column for column in columns if column["pk"]]
    if len(keys) != 1 or keys[0]["type"].upper() != "INTEGER":
        raise ValueError("POI requires one INTEGER feature primary key")
    if "source_fid" in names and keys[0]["name"] != "source_fid":
        raise ValueError("POI source_fid column conflicts with the feature primary key")
    geometry_rows = list(connection.execute("SELECT * FROM gpkg_geometry_columns WHERE table_name=?", (layer,)))
    if len(geometry_rows) != 1 or geometry_rows[0]["column_name"] not in names:
        raise ValueError("POI requires exactly one declared geometry column")
    geometry = dict(geometry_rows[0])
    if geometry["column_name"] in fields.values() or keys[0]["name"] in fields.values():
        raise ValueError("POI attributes cannot alias geometry or feature primary key")
    selected = next(row for row in contents if row["table_name"] == layer)
    if geometry["srs_id"] != selected["srs_id"]:
        raise ValueError("GPKG content and geometry CRS IDs disagree")
    srs_rows = list(connection.execute("SELECT * FROM gpkg_spatial_ref_sys WHERE srs_id=?", (geometry["srs_id"],)))
    if len(srs_rows) != 1:
        raise ValueError("POI CRS definition is missing or ambiguous")
    srs = dict(srs_rows[0])
    crs = CRS.from_wkt(srs["definition"])
    if not crs.equals(CRS.from_epsg(27700)) or not crs.is_projected or any(axis.unit_conversion_factor != 1 for axis in crs.axis_info):
        raise ValueError("POI CRS must be equivalent to EPSG:27700 with metre units")
    count = connection.execute(f"SELECT COUNT(*) FROM {_quoted(layer)}").fetchone()[0]
    cached_count = None
    if connection.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='gpkg_ogr_contents'").fetchone():
        cached = connection.execute("SELECT feature_count FROM gpkg_ogr_contents WHERE table_name=?", (layer,)).fetchone()
        cached_count = cached[0] if cached else None
    return {
        "layers": contents, "layer": layer, "fields": columns, "field_mapping": fields,
        "fid_column": keys[0]["name"], "geometry": geometry, "crs": srs,
        "crs_epsg": crs.to_epsg(), "coordinate_units": "metre",
        "source_count": count, "cached_feature_count": cached_count,
        "cached_count_matches": cached_count == count if cached_count is not None else None,
        "extent_is_coverage_boundary": False,
    }


def poi_batches(relative_path, *, data_root, layer=None, batch_size=10000, fields=None, limit=None):
    """Yield raw GPKG dictionaries in source-FID order, without dropping rows.

    source_fid is the real feature primary key. Geometry remains a GPKG blob;
    malformed geometry must not prevent auditing that row's ID and attributes.
    limit is explicitly a local sample, never a global duplicate audit.
    """
    _positive_integer(batch_size, "batch_size")
    if limit is not None:
        _positive_integer(limit, "limit")
    mapping = _field_mapping(fields, POI_FIELDS)
    with _gpkg_connection(relative_path, data_root) as connection:
        schema = _gpkg_schema(connection, layer, mapping)
        statement = f"SELECT * FROM {_quoted(schema['layer'])} ORDER BY {_quoted(schema['fid_column'])}"
        cursor = connection.execute(statement + (" LIMIT ?" if limit is not None else ""), (limit,) if limit is not None else ())
        while rows := cursor.fetchmany(batch_size):
            batch = []
            for row in rows:
                record = dict(row)
                fid = record[schema["fid_column"]]
                if type(fid) is not int:
                    raise ValueError("feature primary key is not an integer")
                record["source_fid"] = fid
                batch.append(record)
            yield batch


def _raw(value):
    """Lossless typed scalar representation for SQL keys and CSV/JSON output."""
    if value is None:
        return "null", ""
    if isinstance(value, str):
        return "text", value
    if isinstance(value, bytes):
        return "blob", value.hex()
    if type(value) is int:
        return "integer", str(value)
    if type(value) is float:
        return "real", repr(value)
    raise ValueError(f"unsupported raw scalar type: {type(value).__name__}")


def _raw_key(value):
    return json.dumps(_raw(value), ensure_ascii=False, separators=(",", ":"))


def _missing_kind(value):
    if value is None:
        return "null"
    if isinstance(value, str):
        if value == "":
            return "empty_string"
        if not value.strip():
            return "whitespace"
    if isinstance(value, float) and not math.isfinite(value):
        return "nonfinite"
    return "present"


def _classification_parts(value):
    # Same format semantics as classify(); dictionary membership is NOT checked.
    if value is None or (isinstance(value, float) and math.isnan(value)) or value == "":
        return "missing", "UNKNOWN", "UNKNOWN", "UNKNOWN"
    code = str(value)
    if re.fullmatch(r"[0-9]{8}", code) is None:
        return "malformed", "UNKNOWN", "UNKNOWN", "UNKNOWN"
    return "valid", code[:2], code[:4], code


def _number(value):
    if value is None or isinstance(value, (bytes, bool)) or (isinstance(value, str) and not value.strip()):
        return None
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError, OverflowError):
        return None


def _geometry_check(blob, srs_id):
    from shapely import from_wkb
    from shapely.errors import GEOSException

    result = {"status": "null", "type": None, "xy": None, "detail": "", "has_z": False, "has_m": False}
    if blob is None:
        return result
    try:
        if not isinstance(blob, bytes) or len(blob) < 8 or blob[:3] != b"GP\x00":
            raise ValueError("not a version-0 GeoPackage geometry")
        flags = blob[3]
        envelope = (flags >> 1) & 7
        if flags & 0xE0 or envelope not in (0, 1, 2, 3, 4):
            raise ValueError("unsupported GeoPackage flags/envelope")
        header_srs = struct.unpack(("<" if flags & 1 else ">") + "i", blob[4:8])[0]
        if header_srs != srs_id:
            raise ValueError("geometry header CRS differs from layer CRS")
        offset = 8 + (0, 32, 48, 48, 64)[envelope]
        if len(blob) <= offset:
            raise ValueError("truncated GeoPackage geometry")
        geometry = from_wkb(blob[offset:])
        result.update(type=geometry.geom_type, has_z=bool(geometry.has_z), has_m=bool(geometry.has_m))
        if bool(flags & 0x10) != bool(geometry.is_empty):
            raise ValueError("GeoPackage empty flag disagrees with WKB")
        if geometry.is_empty:
            result["status"] = "empty"
        elif geometry.geom_type == "Point" and not all(math.isfinite(v) for v in (geometry.x, geometry.y)):
            result["status"] = "nonfinite"
        elif not geometry.is_valid:
            result["status"] = "invalid"
        elif geometry.geom_type != "Point":
            result["status"] = "non_point"
        else:
            result.update(status="valid_point", xy=(geometry.x, geometry.y))
    except (ValueError, TypeError, GEOSException, struct.error) as error:
        result.update(status="decode_error", detail=str(error), xy=None)
    return result


def _json_write(path, value):
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    temporary.replace(path)


def _csv_writer(stack, path, header):
    stream = stack.enter_context(path.open("x", encoding="utf-8", newline=""))
    writer = csv.writer(stream)
    writer.writerow(header)
    return writer


def _audit_database(path):
    db = sqlite3.connect(path)
    db.execute("PRAGMA cache_size=-8192")
    db.executescript('''
        CREATE TABLE ref_rows (source_fid INTEGER PRIMARY KEY, ref_key TEXT);
        CREATE INDEX ref_key_index ON ref_rows(ref_key);
        CREATE TABLE frequencies (kind TEXT, raw_key TEXT, raw_type TEXT, raw_value TEXT,
            format_status TEXT, group_code TEXT, category_code TEXT, class_code TEXT,
            n INTEGER NOT NULL, PRIMARY KEY(kind, raw_key));
        CREATE TABLE station_rows (record_no INTEGER PRIMARY KEY, station_id TEXT, xy_key TEXT);
        CREATE INDEX station_id_index ON station_rows(station_id);
        CREATE INDEX station_xy_index ON station_rows(xy_key);
    ''')
    return db


def _inspect_pois(path, root, output, db, report, fields, batch_size, limit, known_accuracy, progress):
    info = report["poi"]
    missing = {name: Counter() for name in fields.values()}
    geometry_status, geometry_types, issues = Counter(), Counter(), Counter()
    compared, unequal, mean, maximum, maximum_fid = 0, 0, 0.0, None, None
    with ExitStack() as stack:
        issue_writer = _csv_writer(stack, output / "poi_quality_issues.csv", ["source_fid", "reason", "detail"])
        difference_writer = _csv_writer(stack, output / "coordinate_differences.csv", ["source_fid", "x_geometry", "y_geometry", "feature_easting", "feature_northing", "dx", "dy", "delta_m"])
        samples = stack.enter_context((output / "poi_sample.jsonl").open("x", encoding="utf-8"))
        for batch in poi_batches(path, data_root=root, layer=info["layer"], fields=fields, batch_size=batch_size, limit=limit):
            refs, frequency = [], Counter()
            frequency_values = {}
            for row in batch:
                fid = row["source_fid"]
                for name in fields.values():
                    missing[name][_missing_kind(row[name])] += 1
                ref = row[fields["ref_no"]]
                kind = _missing_kind(ref)
                absent = kind in ("null", "empty_string", "nonfinite")
                refs.append((fid, None if absent else _raw_key(ref)))
                reasons = []
                if absent:
                    reasons.append(("missing_ref_no", kind))
                elif kind == "whitespace":
                    reasons.append(("whitespace_ref_no", _raw_key(ref)))
                for category, field in (("classification", "raw_code"), ("accuracy", "accuracy")):
                    raw_value = row[fields[field]]
                    key = _raw_key(raw_value)
                    frequency[category, key] += 1
                    frequency_values[category, key] = raw_value
                code_status = _classification_parts(row[fields["raw_code"]])[0]
                if code_status != "valid":
                    reasons.append((code_status + "_classification", _raw_key(row[fields["raw_code"]])))
                accuracy = row[fields["accuracy"]]
                if _missing_kind(accuracy) != "present":
                    reasons.append(("missing_or_invalid_accuracy", _raw_key(accuracy)))
                elif known_accuracy is not None and _raw_key(accuracy) not in known_accuracy:
                    reasons.append(("unknown_accuracy_code", _raw_key(accuracy)))
                geo = _geometry_check(row[info["geometry"]["column_name"]], info["geometry"]["srs_id"])
                geometry_status[geo["status"]] += 1
                geometry_types[geo["type"] or "UNREADABLE_OR_NULL"] += 1
                if geo["status"] != "valid_point":
                    reasons.append(("geometry_" + geo["status"], geo["detail"]))
                easting, northing = _number(row[fields["easting"]]), _number(row[fields["northing"]])
                if easting is None or northing is None:
                    reasons.append(("invalid_attribute_coordinates", "finite easting and northing required"))
                delta = None
                if geo["xy"] is not None and easting is not None and northing is not None:
                    x, y = geo["xy"]
                    dx, dy = x - easting, y - northing
                    value = math.hypot(dx, dy)
                    if math.isfinite(value):
                        delta = value
                        compared += 1
                        mean += (delta - mean) / compared
                        if maximum is None or delta > maximum:
                            maximum, maximum_fid = delta, fid
                        difference_writer.writerow([fid, x, y, easting, northing, dx, dy, delta])
                        if delta != 0:
                            unequal += 1
                            reasons.append(("geometry_attribute_difference", repr(delta)))
                    else:
                        reasons.append(("coordinate_difference_overflow", "non-finite difference"))
                for reason, detail in reasons:
                    issues[reason] += 1
                    issue_writer.writerow([fid, reason, detail])
                if info["scanned_count"] < 20:
                    sample = {"scope": "sample", "source_fid": fid, "source_layer": info["layer"],
                              "raw_fields": {name: {"type": _raw(value)[0], "value": _raw(value)[1]} for name, value in row.items() if name != info["geometry"]["column_name"]},
                              "geometry": geo, "delta_m": delta, "issues": reasons}
                    samples.write(json.dumps(sample, ensure_ascii=False, allow_nan=False) + "\n")
                info["scanned_count"] += 1
            db.executemany("INSERT INTO ref_rows VALUES (?,?)", refs)
            entries = []
            for (category, key), count in frequency.items():
                value = frequency_values[category, key]
                raw_type, raw_value = _raw(value)
                parts = _classification_parts(value) if category == "classification" else (None, None, None, None)
                entries.append((category, key, raw_type, raw_value, *parts, count))
            db.executemany("INSERT INTO frequencies VALUES (?,?,?,?,?,?,?,?,?) ON CONFLICT(kind,raw_key) DO UPDATE SET n=n+excluded.n", entries)
            db.commit()
            if progress is not None:
                progress({"stage": "poi", "scanned_count": info["scanned_count"], "source_count": info["source_count"]})
    info.update(missing_by_field={k: dict(v) for k, v in missing.items()}, geometry_status=dict(geometry_status), geometry_types=dict(geometry_types), issue_counts=dict(issues))
    info["coordinate_difference"] = {"comparable_count": compared, "uncomparable_count": info["scanned_count"] - compared, "zero_count": compared - unequal, "nonzero_count": unequal, "mean_m": mean if compared else None, "max_m": maximum, "max_source_fid": maximum_fid}
    expected = info["source_count"] if limit is None else min(info["source_count"], limit)
    info["expected_scan_count"] = expected
    if info["scanned_count"] != expected:
        raise ValueError(f"POI row count mismatch: expected {expected}, scanned {info['scanned_count']}")
    if any(sum(counts.values()) != expected for counts in missing.values()):
        raise ValueError("POI field missingness counts do not reconcile")


def _station_reader(stream, fields):
    text = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
    reader = csv.reader(text, strict=True)
    header = next(reader, None)
    if not header or len(set(header)) != len(header) or any(not name for name in header):
        raise ValueError("station CSV requires a non-empty unique header")
    if set(fields.values()) - set(header):
        raise ValueError(f"station CSV missing required fields: {sorted(set(fields.values()) - set(header))}")
    return text, reader, header


def _inspect_stations(path, root, output, db, report, fields, batch_size, limit):
    info = report["stations"]
    missing = {name: Counter() for name in fields.values()}
    issues = Counter()
    with ExitStack() as stack:
        writer = _csv_writer(stack, output / "station_quality_issues.csv", ["source_record_no", "StationID", "reason", "detail"])
        sample = stack.enter_context((output / "station_sample.jsonl").open("x", encoding="utf-8"))
        stream = stack.enter_context(open_source(path, data_root=root))
        text, reader, header = _station_reader(stream, fields)
        stack.callback(text.detach)
        info.update(header=header, field_mapping=fields)
        iterator = islice(reader, limit) if limit is not None else reader
        for batch in iter(lambda: list(islice(iterator, batch_size)), []):
            entries = []
            for values in batch:
                info["scanned_count"] += 1
                record_no = info["scanned_count"]
                row = {name: values[i] if i < len(values) else None for i, name in enumerate(header)}
                reasons = []
                if len(values) != len(header):
                    reasons.append(("csv_field_count", f"expected {len(header)}, got {len(values)}"))
                for name in fields.values():
                    missing[name][_missing_kind(row[name])] += 1
                station_id = row[fields["id"]]
                if station_id is None or station_id == "":
                    reasons.append(("missing_station_id", "null or empty ID"))
                    station_id = None
                elif not station_id.strip():
                    reasons.append(("whitespace_station_id", repr(station_id)))
                lat, lon = _number(row[fields["latitude"]]), _number(row[fields["longitude"]])
                xy_key = None
                if lat is None or lon is None:
                    reasons.append(("invalid_station_coordinates", "finite numeric latitude/longitude required"))
                elif not (-90 <= lat <= 90 and -180 <= lon <= 180):
                    reasons.append(("station_coordinates_out_of_range", f"longitude={lon}, latitude={lat}"))
                else:
                    # Signed zero has the same exact numerical location.
                    xy_key = json.dumps([lon or 0.0, lat or 0.0], separators=(",", ":"))
                entries.append((record_no, station_id, xy_key))
                for reason, detail in reasons:
                    writer.writerow([record_no, station_id, reason, detail])
                    issues[reason] += 1
                if record_no <= 20:
                    sample.write(json.dumps({"scope": "sample", "source_record_no": record_no, "raw_fields": row, "issues": reasons}, ensure_ascii=False) + "\n")
            db.executemany("INSERT INTO station_rows VALUES (?,?,?)", entries)
            db.commit()
        group_checks = [
            ("duplicate_station_id", "station_id", "station_id IS NOT NULL", "COUNT(*) > 1"),
            ("station_id_position_conflict", "station_id", "station_id IS NOT NULL AND xy_key IS NOT NULL", "COUNT(DISTINCT xy_key) > 1"),
            ("co_located_different_ids", "xy_key", "station_id IS NOT NULL AND xy_key IS NOT NULL", "COUNT(DISTINCT station_id) > 1"),
        ]
        groups = {}
        for reason, key, where, having in group_checks:
            group_count, affected = 0, 0
            query = f"SELECT {key} FROM station_rows WHERE {where} GROUP BY {key} HAVING {having} ORDER BY {key}"
            for (value,) in db.execute(query):
                group_count += 1
                for record_no, station_id in db.execute(f"SELECT record_no,station_id FROM station_rows WHERE {key}=? AND station_id IS NOT NULL ORDER BY record_no", (value,)):
                    writer.writerow([record_no, station_id, reason, value])
                    affected += 1
                    issues[reason] += 1
            groups[reason] = {"group_count": group_count, "affected_row_count": affected}
    info.update(missing_by_field={k: dict(v) for k, v in missing.items()}, issue_counts=dict(issues), groups=groups, source_count=None)
    if limit is None:
        # Independent logical-record count: quoted newlines are not records.
        with open_source(path, data_root=root) as stream:
            text = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
            try:
                reader = csv.reader(text, strict=True)
                next(reader)
                info["source_count"] = sum(1 for _ in reader)
            finally:
                text.detach()
        if info["source_count"] != info["scanned_count"]:
            raise ValueError("station logical record counts do not reconcile")


def _export_frequencies(db, output, count, known_accuracy):
    summaries = {}
    with ExitStack() as stack:
        for kind, filename in (("classification", "classification_frequencies.csv"), ("accuracy", "positional_accuracy_frequencies.csv")):
            writer = _csv_writer(stack, output / filename, ["raw_type", "raw_value", "format_status", "group_code", "category_code", "class_code", "count", "dictionary_status", "accuracy_description"])
            total, distinct = 0, 0
            for key, raw_type, raw_value, status, group, category, class_code, n in db.execute("SELECT raw_key,raw_type,raw_value,format_status,group_code,category_code,class_code,n FROM frequencies WHERE kind=? ORDER BY raw_key", (kind,)):
                writer.writerow([raw_type, raw_value, status, group, category, class_code, n, "not_checked" if kind == "classification" else "", (known_accuracy or {}).get(key, "") if kind == "accuracy" else ""])
                total += n
                distinct += 1
            if total != count:
                raise ValueError(f"{kind} frequency count {total} differs from scanned count {count}")
            summaries[kind] = {"row_count": total, "distinct_raw_values": distinct}
        writer = _csv_writer(stack, output / "native_level_frequencies.csv", ["level", "code", "count"])
        for level in ("group", "category", "class"):
            counts = db.execute(f"SELECT {level}_code,SUM(n) FROM frequencies WHERE kind='classification' GROUP BY {level}_code ORDER BY {level}_code")
            total, distinct = 0, 0
            for code, n in counts:
                writer.writerow([level, code, n])
                total += n
                distinct += 1
            if total != count:
                raise ValueError(f"{level} counts do not reconcile")
            summaries[level] = {"row_count": total, "distinct_codes_including_unknown": distinct}
    return summaries


def _export_duplicates(db, output, source, layer, sample):
    db.execute("CREATE TABLE duplicate_refs (ref_key TEXT PRIMARY KEY, n INTEGER)")
    db.execute("INSERT INTO duplicate_refs SELECT ref_key,COUNT(*) FROM ref_rows WHERE ref_key IS NOT NULL GROUP BY ref_key HAVING COUNT(*)>1")
    db.commit()
    prefix = "sample_" if sample else ""
    key_count, involved = 0, 0
    with ExitStack() as stack:
        keys = _csv_writer(stack, output / (prefix + "duplicate_ref_nos.csv"), ["raw_type", "raw_value", "count"])
        rows = _csv_writer(stack, output / (prefix + "duplicate_ref_rows.csv"), ["source_path", "source_layer", "source_fid", "raw_type", "raw_value", "count"])
        for key, n in db.execute("SELECT ref_key,n FROM duplicate_refs ORDER BY ref_key"):
            raw_type, raw_value = json.loads(key)
            keys.writerow([raw_type, raw_value, n])
            key_count += 1
            involved += n
            actual = 0
            for (fid,) in db.execute("SELECT source_fid FROM ref_rows WHERE ref_key=? ORDER BY source_fid", (key,)):
                rows.writerow([source, layer, fid, raw_type, raw_value, n])
                actual += 1
            if actual != n:
                raise ValueError("duplicate source rows do not reconcile")
    missing = db.execute("SELECT COUNT(*) FROM ref_rows WHERE ref_key IS NULL").fetchone()[0]
    unique = db.execute("SELECT COUNT(DISTINCT ref_key) FROM ref_rows").fetchone()[0]
    total = db.execute("SELECT COUNT(*) FROM ref_rows").fetchone()[0]
    if missing + unique + involved - key_count != total:
        raise ValueError("ref_no counts do not reconcile")
    return {"missing_count": missing, "unique_nonmissing_count": unique, "duplicate_key_count": key_count, "duplicate_row_count": involved, "excess_row_count": involved - key_count, "source_fid_count": total}


def inspect_inputs(poi_path, station_path, *, data_root, output_dir, layer=None,
                   poi_batch_size=10000, station_batch_size=1000, sample_limit=None,
                   poi_fields=None, station_fields=None, accuracy_labels=None, progress=None):
    """Audit raw inputs into a NEW outputs/<run-id> directory and return a report.

    Caller verifies the source manifest first. This function additionally hashes
    its two inputs before/after reading. sample_limit never produces a complete
    duplicate list. Errors retain a failed report and raise InputInspectionError.
    """
    _positive_integer(poi_batch_size, "poi_batch_size")
    _positive_integer(station_batch_size, "station_batch_size")
    if sample_limit is not None:
        _positive_integer(sample_limit, "sample_limit")
    fields = _field_mapping(poi_fields, POI_FIELDS)
    stations = _field_mapping(station_fields, STATION_FIELDS)
    labels = None
    if accuracy_labels is not None:
        if not isinstance(accuracy_labels, dict) or any(not isinstance(v, str) for v in accuracy_labels.values()):
            raise ValueError("accuracy_labels must map raw codes to descriptions")
        labels = {_raw_key(k): v for k, v in accuracy_labels.items()}
    root = project_path(".", root=data_root)
    output = project_path(output_dir, root=root, field="output_dir")
    if not Path(output_dir).is_relative_to("outputs") or output == root / "outputs" or not output.is_relative_to(root / "outputs"):
        raise ValueError("output_dir must be a new directory strictly inside outputs/")
    for source in (poi_path, station_path):
        _source_path(source, root)
    output.mkdir(parents=True, exist_ok=False)
    report = {
        "schema_version": 1, "status": "running", "scope": "sample" if sample_limit is not None else "full_source",
        "sample_limit": sample_limit, "scan_complete": False, "full_source_complete": False,
        "duplicate_list_complete": False, "errors": [], "output_dir": str(output),
        "poi": {"path": poi_path, "scanned_count": 0}, "stations": {"path": station_path, "scanned_count": 0},
        "batch_sizes": {"poi": poi_batch_size, "stations": station_batch_size},
        "dictionary_checked": False, "selection_or_cleaning_applied": False,
        "accuracy_lookup_supplied": labels is not None,
    }
    _json_write(output / "input_inspection.json", report)
    db = None
    try:
        before = {source: _fingerprint(source, root, 1048576) for source in (poi_path, station_path)}
        states = {source: _file_state(_source_path(source, root).stat()) for source in before}
        report["input_fingerprints_before"] = before
        with _gpkg_connection(poi_path, root) as connection:
            report["poi"].update(_gpkg_schema(connection, layer, fields))
        db = _audit_database(output / "audit_index.sqlite")
        _inspect_pois(poi_path, root, output, db, report, fields, poi_batch_size, sample_limit, labels, progress)
        _inspect_stations(station_path, root, output, db, report, stations, station_batch_size, sample_limit)
        report["poi"]["frequencies"] = _export_frequencies(db, output, report["poi"]["scanned_count"], labels)
        after = {source: _fingerprint(source, root, 1048576) for source in before}
        report["input_fingerprints_after"] = after
        report["input_bytes_unchanged"] = before == after
        if before != after or any(states[source] != _file_state(_source_path(source, root).stat()) for source in before):
            raise ValueError("input changed during inspection")
        report["poi"]["ref_no"] = _export_duplicates(db, output, poi_path, report["poi"]["layer"], sample_limit is not None)
        if report["poi"]["ref_no"]["source_fid_count"] != report["poi"]["scanned_count"]:
            raise ValueError("source FID count does not reconcile")
        report.update(status="completed", scan_complete=True, full_source_complete=sample_limit is None, duplicate_list_complete=sample_limit is None)
        report["quality_issues_present"] = bool(report["poi"]["issue_counts"] or report["stations"]["issue_counts"] or report["poi"]["ref_no"]["duplicate_key_count"])
    except Exception as error:
        report.update(status="failed", scan_complete=False, full_source_complete=False, duplicate_list_complete=False)
        report["errors"].append(f"{type(error).__name__}: {error}")
        _json_write(output / "input_inspection.json", report)
        raise InputInspectionError(report) from error
    finally:
        if db is not None:
            db.close()
    _json_write(output / "input_inspection.json", report)
    return report


class DictionaryValidationError(ValueError):
    """Failed step-07 report; nearby partial files are not usable results."""

    def __init__(self, report):
        self.report = report
        super().__init__("; ".join(report["errors"]))


def _artifact_fingerprint(path):
    digest, size = hashlib.sha256(), 0
    with _open_regular(path) as stream:
        before = _file_state(os.fstat(stream.fileno()))
        while block := stream.read(1048576):
            digest.update(block)
            size += len(block)
        if before != _file_state(os.fstat(stream.fileno())) or before != _file_state(path.stat()):
            raise ValueError(f"audit artifact changed: {path}")
    return {"sha256": digest.hexdigest(), "size_bytes": size}


def _audit_path(relative, root):
    path = project_path(relative, root=root)
    if not Path(relative).is_relative_to("outputs") or not path.is_relative_to(root / "outputs"):
        raise ValueError("audit artifacts must stay inside outputs/")
    return path


def _read_frequency_csv(path):
    with _open_regular(path) as stream:
        text = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
        try:
            reader = csv.DictReader(text, strict=True)
            if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
                raise ValueError("invalid frequency header")
            result = []
            for row in reader:
                if None in row or any(v is None for v in row.values()):
                    raise ValueError("invalid frequency field count")
                if re.fullmatch(r"[1-9][0-9]*", row["count"]) is None:
                    raise ValueError("frequency must be a positive integer")
                row["count"] = int(row["count"])
                result.append(row)
            return result
        finally:
            text.detach()


def _frequency_raw(row):
    kind, value = row["raw_type"], row["raw_value"]
    if kind == "text":
        raw = value
    elif kind == "null" and value == "":
        raw = None
    elif kind == "integer":
        raw = int(value)
    elif kind == "real":
        raw = float(value)
    elif kind == "blob":
        raw = bytes.fromhex(value)
    else:
        raise ValueError("invalid raw classification representation")
    if _raw(raw) != (kind, value):
        raise ValueError("noncanonical raw classification representation")
    return raw


def load_inspection_frequencies(inspection_dir, *, data_root, poi_path, layer):
    """Require complete step-06 evidence and byte-identical current POI input."""
    root = project_path(".", root=data_root)
    directory = _audit_path(inspection_dir, root)
    evidence_path = _audit_path(str((directory.parent / "verification.json").relative_to(root)), root)
    with _open_regular(evidence_path) as stream:
        evidence = json.load(stream, object_pairs_hook=_unique_json_object)
    if evidence.get("status") != "passed":
        raise ValueError("prior verification is not passed")
    fingerprints = {}
    for name in ("input_inspection.json", "classification_frequencies.csv", "native_level_frequencies.csv"):
        path = _audit_path(str((directory / name).relative_to(root)), root)
        actual = _artifact_fingerprint(path)
        expected = evidence["artifacts"].get(str(path.relative_to(evidence_path.parent)))
        if actual != expected:
            raise ValueError(f"prior artifact checksum mismatch: {name}")
        fingerprints[name] = actual
    with _open_regular(directory / "input_inspection.json") as stream:
        report = json.load(stream, object_pairs_hook=_unique_json_object)
    if (report.get("status") != "completed" or report.get("scope") != "full_source"
            or any(report.get(flag) is not True for flag in ("scan_complete", "full_source_complete", "duplicate_list_complete", "input_bytes_unchanged"))):
        raise ValueError("prior audit must be completed full_source, never sample/failed")
    if report["poi"]["path"] != poi_path or report["poi"]["layer"] != layer:
        raise ValueError("prior POI path or layer differs")
    current = _fingerprint(poi_path, root, 1048576)
    if current != report["input_fingerprints_after"].get(poi_path) or current != report["input_fingerprints_before"].get(poi_path):
        raise ValueError("current POI fingerprint differs from prior full audit")
    rows = _read_frequency_csv(directory / "classification_frequencies.csv")
    levels, raw_seen = {level: Counter() for level in ("group", "category", "class")}, set()
    for row in rows:
        key = row["raw_type"], row["raw_value"]
        if key in raw_seen:
            raise ValueError("duplicate raw frequency key")
        raw_seen.add(key)
        parts = _classification_parts(_frequency_raw(row))
        if parts != tuple(row[name] for name in ("format_status", "group_code", "category_code", "class_code")):
            raise ValueError("prior raw classification format or prefixes disagree")
        for level in levels:
            levels[level][row[f"{level}_code"]] += row["count"]
    saved_levels = {level: {} for level in levels}
    for row in _read_frequency_csv(directory / "native_level_frequencies.csv"):
        level, code = row["level"], row["code"]
        if level not in saved_levels or code in saved_levels[level]:
            raise ValueError("invalid or duplicate level frequency key")
        saved_levels[level][code] = row["count"]
    count = report["poi"]["scanned_count"]
    if type(count) is not int or count < 0 or count != report["poi"]["source_count"]:
        raise ValueError("prior source counts disagree")
    if saved_levels != levels or any(sum(values.values()) != count for values in levels.values()):
        raise ValueError("prior frequency totals disagree")
    for name, before in fingerprints.items():
        if before != _artifact_fingerprint(directory / name):
            raise ValueError("prior artifact changed while reading")
    return {"rows": rows, "level_counts": levels, "source_count": count,
            "artifacts": fingerprints, "verification": _artifact_fingerprint(evidence_path),
            "poi_fingerprint": current}


def source_name_frequencies(poi_path, *, data_root, layer, index_path, batch_size=10000, expected_rows, progress=None):
    """Scan only code and category-name attributes; no geometry decoding."""
    _positive_integer(batch_size, "batch_size")
    root = project_path(".", root=data_root)
    target = _audit_path(index_path, root)
    if target.exists():
        raise FileExistsError(target)
    with sqlite3.connect(target) as index:
        index.execute("PRAGMA cache_size=-8192")
        index.execute("CREATE TABLE names (level TEXT, code TEXT, raw_key TEXT, raw_type TEXT, raw_value TEXT, n INTEGER, PRIMARY KEY(level,code,raw_key))")
        index.execute("CREATE TABLE codes (raw_key TEXT PRIMARY KEY, n INTEGER)")
        count = 0
        with _gpkg_connection(poi_path, root) as connection:
            schema = _gpkg_schema(connection, layer, POI_FIELDS)
            columns = ("pointx_class", "groupname", "categoryname", "classname")
            if not set(columns).issubset({row["name"] for row in schema["fields"]}):
                raise ValueError("POI classification name fields are missing")
            cursor = connection.execute(f"SELECT {','.join(_quoted(c) for c in columns)} FROM {_quoted(layer)} ORDER BY {_quoted(schema['fid_column'])}")
            for batch in iter(lambda: cursor.fetchmany(batch_size), []):
                names, codes = Counter(), Counter()
                for row in batch:
                    codes[_raw_key(row[0])] += 1
                    parts = _classification_parts(row[0])
                    for i, level in enumerate(("group", "category", "class"), 1):
                        if row[i] is not None and not isinstance(row[i], str):
                            raise ValueError("source classification names must be text or NULL")
                        raw_type, raw_value = _raw(row[i])
                        names[level, parts[i], _raw_key(row[i]), raw_type, raw_value] += 1
                index.executemany("INSERT INTO names VALUES (?,?,?,?,?,?) ON CONFLICT(level,code,raw_key) DO UPDATE SET n=n+excluded.n", [(*key, n) for key, n in names.items()])
                index.executemany("INSERT INTO codes VALUES (?,?) ON CONFLICT(raw_key) DO UPDATE SET n=n+excluded.n", list(codes.items()))
                index.commit()
                count += len(batch)
                if progress:
                    progress({"stage": "source_names", "scanned_count": count, "source_count": schema["source_count"]})
            if count != schema["source_count"]:
                raise ValueError("source-name scan incomplete")
        expected = {_raw_key(_frequency_raw(row)): row["count"] for row in expected_rows}
        if dict(index.execute("SELECT raw_key,n FROM codes")) != expected:
            raise ValueError("source-name scan raw frequencies differ from prior audit")
        return [{"level": level, "code": code, "raw_type": raw_type, "raw_value": raw_value, "count": n}
                for level, code, raw_type, raw_value, n in index.execute("SELECT level,code,raw_type,raw_value,n FROM names ORDER BY level,code,raw_key")]


def _write_records(path, rows, header):
    with path.open("x", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=header, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def inspect_dictionary(group_path, category_path, class_path, poi_path, *, data_root,
                       inspection_dir, output_dir, layer, required_paths,
                       batch_size=10000, progress=None):
    """Build step-07 dictionary/labels with source and prior-artifact verification.

    This is an explicitly called step helper, not a full pipeline or CLI.
    Consumers must require status=completed and validation_complete=true.
    """
    from src.classification import native_dictionary, merge_classification_frequencies
    from src.analysis import category_labels

    _positive_integer(batch_size, "batch_size")
    root = project_path(".", root=data_root)
    output = _audit_path(output_dir, root)
    if output == root / "outputs":
        raise ValueError("output must be strictly inside outputs/")
    output.mkdir(parents=True, exist_ok=False)
    report = {"schema_version": 1, "status": "running", "validation_complete": False,
              "errors": [], "output_dir": str(output), "batch_size": batch_size,
              "prior_inspection_dir": inspection_dir, "source_scope": "full_source",
              "label_rule": "source frequency descending, Unicode name ascending; dictionary then code fallback",
              "selection_or_cleaning_applied": False}
    _json_write(output / "dictionary_validation.json", report)
    try:
        paths = (group_path, category_path, class_path, poi_path)
        if not set(paths).issubset(set(required_paths)):
            raise ValueError("required_paths must include all dictionary and POI inputs")
        report["sources_before"] = verify_sources(data_root=root, required_paths=required_paths)
        before = {path: _fingerprint(path, root, 1048576) for path in paths}
        states = {path: _file_state(_source_path(path, root).stat()) for path in paths}
        report["input_fingerprints_before"] = before
        prior = load_inspection_frequencies(inspection_dir, data_root=root, poi_path=poi_path, layer=layer)
        report["prior_artifacts"] = prior["artifacts"]
        report["prior_verification"] = prior["verification"]
        dictionary = native_dictionary(group_path, category_path, class_path, data_root=root)
        merged = merge_classification_frequencies(dictionary, prior["rows"])
        names = source_name_frequencies(poi_path, data_root=root, layer=layer,
                    index_path=str((output / "name_index.sqlite").relative_to(root)),
                    batch_size=batch_size, expected_rows=prior["rows"], progress=progress)
        labels = category_labels(dictionary, merged["level_counts"], names)
        report["source_count"] = merged["source_count"]
        report["dictionary_counts"] = {level: len(table) for level, table in (("group", dictionary.groups), ("category", dictionary.categories), ("class", dictionary.classes))}
        report["levels"] = {level: {"source_count": sum(row["source_count"] for row in labels if row["level"] == level),
                          "matching": dict(Counter(row["dictionary_status"] for row in labels if row["level"] == level)),
                          "label_sources": dict(Counter(row["label_source"] for row in labels if row["level"] == level)),
                          "name_conflict_codes": sum(row["source_name_conflict"] for row in labels if row["level"] == level),
                          "missing_name_rows": sum(row["source_name_missing_count"] for row in labels if row["level"] == level),
                          "source_dictionary_name_differences": sum(row["source_dictionary_name_difference"] for row in labels if row["level"] == level)}
                          for level in merged["level_counts"]}
        report["sources_after"] = verify_sources(data_root=root, required_paths=required_paths)
        after = {path: _fingerprint(path, root, 1048576) for path in paths}
        report["input_fingerprints_after"] = after
        if before != after or any(states[path] != _file_state(_source_path(path, root).stat()) for path in paths):
            raise ValueError("dictionary or POI input changed during processing")
        check_prior = load_inspection_frequencies(inspection_dir, data_root=root, poi_path=poi_path, layer=layer)
        if check_prior != prior:
            raise ValueError("prior audit artifacts changed during processing")
        _write_records(output / "native_dictionary.csv", dictionary.rows(), ["level", "code", "parent_code", "local_code", "foreign_key", "dictionary_name", "source_path", "source_record_no", "source_id"])
        _write_records(output / "classification_dictionary.csv", merged["classes"], ["class_code", "group_code", "category_code", "group_dictionary_name", "category_dictionary_name", "class_dictionary_name", "source_count", "dictionary_status"])
        _write_records(output / "classification_frequency_audit.csv", merged["raw_audit"], ["raw_type", "raw_value", "format_status", "group_code", "category_code", "class_code", "count", "dictionary_status"])
        _write_records(output / "source_name_frequencies.csv", names, ["level", "code", "raw_type", "raw_value", "count"])
        _write_records(output / "category_labels.csv", labels, list(labels[0]) if labels else ["level", "code", "label"])
        report["artifacts"] = {path.name: _artifact_fingerprint(path) for path in sorted(output.iterdir()) if path.suffix in (".csv", ".sqlite")}
        report.update(status="completed", validation_complete=True, input_bytes_unchanged=True)
    except Exception as error:
        report.update(status="failed", validation_complete=False)
        report["errors"].append(f"{type(error).__name__}: {error}")
        _json_write(output / "dictionary_validation.json", report)
        raise DictionaryValidationError(report) from error
    _json_write(output / "dictionary_validation.json", report)
    return report


# Step 11: no station-neighbour queries and no real full-source cleaning here.
@dataclass(frozen=True)
class POICleaningContext:
    source_path: str
    layer: str
    schema: dict
    duplicate_keys: object
    duplicate_rows: object
    coverage: object
    dictionary: object
    rules: object
    accuracy_labels: object
    provenance: dict


def _sealed_step(directory, names, root):
    """Bind each consumed artifact to the parent run's passed verification."""
    directory = _audit_path(directory, root)
    evidence_path = directory.parent / "verification.json"
    evidence = _read_json(evidence_path)
    if evidence.get("status") != "passed":
        raise ValueError(f"upstream verification is not passed: {evidence_path}")
    files = {}
    for name in names:
        path = _audit_path(str((directory / name).relative_to(root)), root)
        actual = _artifact_fingerprint(path)
        if actual != evidence.get("artifacts", {}).get(str(path.relative_to(directory.parent))):
            raise ValueError(f"upstream artifact checksum mismatch: {path}")
        files[name] = actual
    return directory, {"directory": str(directory.relative_to(root)), "artifacts": files,
                       "verification": _artifact_fingerprint(evidence_path)}


def _completed_evidence(report, label, *, full=False):
    if (report.get("status") != "completed" or report.get("validation_complete") is not True
            or report.get("input_bytes_unchanged") is not True
            or (full and report.get("scope") != "full_source")):
        raise ValueError(f"{label} must be completed, validated and byte-stable")


def _check_prior_inputs(before, after, current, label):
    if not before or before != after:
        raise ValueError(f"{label} missing/inconsistent before/after fingerprints")
    for name, fingerprint in before.items():
        if current.get(name) != fingerprint:
            raise ValueError(f"{label} current input fingerprint mismatch: {name}")


def _read_duplicate_evidence(directory, report, source, layer):
    keys, rows, observed = {}, {}, Counter()
    with (directory / "duplicate_ref_nos.csv").open(encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream, strict=True)
        if reader.fieldnames != ["raw_type", "raw_value", "count"]:
            raise ValueError("full-source duplicate key header mismatch")
        for row in reader:
            raw = _frequency_raw(row)
            if _missing_kind(raw) in ("null", "empty_string", "nonfinite"):
                raise ValueError("duplicate evidence contains missing ID")
            key = _raw_key(raw)
            if row.keys() != {"raw_type", "raw_value", "count"} or not re.fullmatch(r"[1-9][0-9]*", row["count"]) or int(row["count"]) < 2 or key in keys:
                raise ValueError("invalid/duplicate full-source duplicate key")
            keys[key] = int(row["count"])
    with (directory / "duplicate_ref_rows.csv").open(encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream, strict=True)
        header = ["source_path", "source_layer", "source_fid", "raw_type", "raw_value", "count"]
        if reader.fieldnames != header:
            raise ValueError("full-source duplicate row header mismatch")
        for row in reader:
            if row.keys() != set(header) or any(v is None for v in row.values()):
                raise ValueError("malformed full-source duplicate row")
            if row["source_path"] != source or row["source_layer"] != layer:
                raise ValueError("duplicate evidence source/layer mismatch")
            fid, key = int(row["source_fid"]), _raw_key(_frequency_raw(row))
            if fid in rows or key not in keys or int(row["count"]) != keys[key]:
                raise ValueError("full-source duplicate row/key/count mismatch")
            rows[fid] = key
            observed[key] += 1
    if observed != keys:
        raise ValueError("full-source duplicate evidence omits conflict records")
    counts = report["poi"]["ref_no"]
    n = report["poi"]["source_count"]
    expected = {"duplicate_key_count": len(keys), "duplicate_row_count": len(rows),
                "excess_row_count": len(rows) - len(keys), "source_fid_count": n}
    if (any(counts.get(k) != v for k, v in expected.items())
            or counts["missing_count"] + counts["unique_nonmissing_count"] + len(rows) - len(keys) != n):
        raise ValueError("full-source duplicate evidence totals disagree")
    return MappingProxyType(keys), MappingProxyType(rows)


def load_poi_cleaning_context(poi_path, *, data_root, code_root, layer,
                              inspection_dir, dictionary_dir, coverage_dir, scope_dir,
                              group_path, category_path, class_path, accuracy_path,
                              rules_config="config/poi_research_rules_v3.yaml", required_paths):
    """Load verified step 06/07/09/10 evidence, without scanning POI geometries.

    Manifest verification hashes all declared input bytes; this is not POI
    cleaning. The complete source conflict list must correspond to this exact
    POI file, regardless of which subset is later passed to clean_poi_batch.
    """
    from pyproj import CRS
    from shapely import from_wkb, prepare
    from shapely.geometry import box
    from src.classification import native_dictionary
    from src.poi_research import load_research_rules

    root = project_path(".", root=data_root)
    direct = {poi_path, group_path, category_path, class_path, accuracy_path}
    if not direct.issubset(set(required_paths)):
        raise ValueError("required_paths must independently include POI/dictionaries/accuracy")
    manifest_path = root / "data/source_manifest.json"
    manifest_fp = _artifact_fingerprint(manifest_path)
    sources = verify_sources(data_root=root, required_paths=required_paths)
    current = {r["path"]: {"resolved_path": r["resolved_path"], "size_bytes": r["actual_size_bytes"], "sha256": r["actual_sha256"]} for r in sources["files"]}
    six, evidence6 = _sealed_step(inspection_dir, ("input_inspection.json", "duplicate_ref_nos.csv", "duplicate_ref_rows.csv"), root)
    report6 = _read_json(six / "input_inspection.json")
    if (report6.get("status") != "completed" or report6.get("scope") != "full_source"
            or report6.get("sample_limit") is not None
            or any(report6.get(k) is not True for k in ("scan_complete", "full_source_complete", "duplicate_list_complete", "input_bytes_unchanged"))):
        raise ValueError("step06 duplicate evidence must be complete full_source, never sample")
    poi = report6["poi"]
    if poi["path"] != poi_path or poi["layer"] != layer or poi["scanned_count"] != poi["source_count"]:
        raise ValueError("step06 POI source/layer/count mismatch")
    _check_prior_inputs(report6["input_fingerprints_before"], report6["input_fingerprints_after"], current, "step06")
    duplicate_keys, duplicate_rows = _read_duplicate_evidence(six, report6, poi_path, layer)
    with _gpkg_connection(poi_path, root) as connection:
        schema = _gpkg_schema(connection, layer, poi["field_mapping"])
    if schema != {k: poi[k] for k in schema}:
        raise ValueError("POI current schema differs from complete step06 evidence")
    seven, evidence7 = _sealed_step(dictionary_dir, ("dictionary_validation.json", "native_dictionary.csv", "classification_dictionary.csv"), root)
    report7 = _read_json(seven / "dictionary_validation.json")
    _completed_evidence(report7, "step07")
    if report7.get("source_scope") != "full_source" or report7.get("prior_inspection_dir") != inspection_dir or report7["source_count"] != poi["source_count"]:
        raise ValueError("step07 dictionary must bind complete step06 source evidence")
    _check_prior_inputs(report7["input_fingerprints_before"], report7["input_fingerprints_after"], current, "step07")
    dictionary = native_dictionary(group_path, category_path, class_path, data_root=root)
    with (seven / "native_dictionary.csv").open(encoding="utf-8", newline="") as stream:
        saved = list(csv.DictReader(stream, strict=True))
    expected = [{k: str(r[k]) for k in saved[0]} for r in dictionary.rows()] if saved else []
    if saved != expected or report7["dictionary_counts"] != {"group": len(dictionary.groups), "category": len(dictionary.categories), "class": len(dictionary.classes)}:
        raise ValueError("step07 dictionary rows/counts differ from current official dictionary")
    nine, evidence9 = _sealed_step(scope_dir, ("pua_scope_report.json", "scope_definition.json", "study_area.wkb"), root)
    report9 = _read_json(nine / "pua_scope_report.json")
    _completed_evidence(report9, "step09")
    if report9.get("scope_complete") is not True:
        raise ValueError("step09 PUA scope is incomplete")
    scope = _read_json(nine / "scope_definition.json")
    _check_prior_inputs(report9["input_fingerprints_before"], report9["input_fingerprints_after"], current, "step09")
    ten, evidence10 = _sealed_step(coverage_dir, ("station_preparation_report.json", "coverage_definition.json", "data_coverage_proxy.wkb", "gb_proxy.wkb"), root)
    report10 = _read_json(ten / "station_preparation_report.json")
    _completed_evidence(report10, "step10", full=True)
    if report10.get("full_source_identity_checked") is not True:
        raise ValueError("step10 lacks full-source identity verification")
    _check_prior_inputs(report10["input_fingerprints_before"], report10["input_fingerprints_after"], current, "step10")
    definition = _read_json(ten / "coverage_definition.json")
    if (definition["input_fingerprints"] != report10["input_fingerprints_before"]
            or definition["target_puas"] != scope["target_puas"]
            or definition["upstream_scope"]["scope_dir"] != scope_dir
            or definition["upstream_scope"]["verification"] != evidence9["verification"]
            or any(definition["upstream_scope"]["artifacts"].get(k) != v for k, v in evidence9["artifacts"].items())):
        raise ValueError("step10 coverage differs from bound step09/source evidence")
    if (definition["predicate"] != "covers" or definition["units"] != "metre"
            or not CRS.from_user_input(definition["crs"]).equals(CRS.from_epsg(27700))):
        raise ValueError("coverage requires EPSG:27700 metre CRS and covers")
    order = definition["order"]
    if order["poi_path"] != poi_path or order["poi_layer"] != layer or order["poi_metadata"] != schema:
        raise ValueError("coverage order POI source/layer/schema mismatch")
    if not set(definition["input_fingerprints"]).issubset(set(required_paths)):
        raise ValueError("required_paths must include all verified coverage inputs")
    coverage, gb, study = (from_wkb(p.read_bytes()) for p in (ten / "data_coverage_proxy.wkb", ten / "gb_proxy.wkb", nine / "study_area.wkb"))
    for label, geom in (("coverage", coverage), ("GB", gb), ("PUA", study)):
        if geom.geom_type not in ("Polygon", "MultiPolygon") or not geom.is_valid or geom.has_z or geom.has_m:
            raise ValueError(f"{label} must be valid 2D Polygon/MultiPolygon")
    if not gb.covers(study) or not coverage.equals(gb.intersection(box(*order["bounds"]))):
        raise ValueError("coverage WKB is not the verified GB/order intersection")
    prepare(coverage)
    with open_source(accuracy_path, data_root=root) as stream:
        text = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
        try:
            reader = csv.reader(text, delimiter="|", strict=True)
            if next(reader) != ["Positional Accuracy Code", "Description"]:
                raise ValueError("accuracy lookup header mismatch")
            labels = {}
            for values in reader:
                if len(values) != 2 or values[0] not in ("1", "2", "3", "4") or values[0] in labels or not values[1]:
                    raise ValueError("accuracy lookup invalid/duplicate row")
                labels[values[0]] = values[1]
            if set(labels) != {"1", "2", "3", "4"}:
                raise ValueError("accuracy lookup incomplete")
        finally:
            text.detach()
    dict_inputs = {p: current[p] for p in (group_path, category_path, class_path)}
    rules = load_research_rules(rules_config, code_root=code_root, dictionary=dictionary, dictionary_inputs=dict_inputs)
    evidence = {"step06": evidence6, "step07": evidence7, "step09": evidence9, "step10": evidence10}
    # Rehash small consumed artifacts, and verify every raw input before returning.
    for item in evidence.values():
        _, checked = _sealed_step(item["directory"], tuple(item["artifacts"]), root)
        if checked != item:
            raise ValueError("upstream evidence changed during context loading")
    if manifest_fp != _artifact_fingerprint(manifest_path):
        raise ValueError("manifest changed during context loading")
    after = verify_sources(data_root=root, required_paths=required_paths)
    if after != sources:
        raise ValueError("sources changed during context loading")
    return POICleaningContext(poi_path, layer, schema, duplicate_keys, duplicate_rows,
                             coverage, dictionary, rules, MappingProxyType(labels),
                             {"upstream": evidence, "input_fingerprints": current, "manifest": manifest_fp,
                              "sources": sources, "coverage_definition": definition, "scope_definition": scope,
                              "rules": rules.provenance})


def _poi_typed(value):
    # bool cannot occur as a distinct SQLite scalar, but is auditable in hand cases.
    if type(value) is bool:
        return {"type": "boolean", "value": "true" if value else "false"}
    kind, raw = _raw(value)
    return {"type": kind, "value": raw}


def clean_poi_batch(records, *, context, include_ev_charging=False):
    """One audit per raw record, all checks independent; input is never mutated.

    Returns base_pois, analysis_pois, audit and summary. No PUA filter, neighbour
    query, spatial deduplication or geometry repair is performed.
    """
    from pyproj import CRS
    from pyproj.exceptions import CRSError
    from shapely.geometry import Point
    from src.classification import classify
    from src.poi_research import research_classification

    if not isinstance(context, POICleaningContext) or type(include_ev_charging) is not bool:
        raise ValueError("verified POICleaningContext and boolean facility switch required")
    schema = context.schema
    crs = CRS.from_wkt(schema["crs"]["definition"])
    if not crs.equals(CRS.from_epsg(27700)) or schema["coordinate_units"] != "metre":
        raise ValueError("POI context CRS must be EPSG:27700 in metres")
    coverage = context.coverage
    if coverage.geom_type not in ("Polygon", "MultiPolygon") or not coverage.is_valid or coverage.has_z or coverage.has_m:
        raise ValueError("POI coverage must be a valid 2D Polygon/MultiPolygon")
    fields, geom_col = schema["field_mapping"], schema["geometry"]["column_name"]
    if set(context.rules.rows) - {"10540832"} != set(context.dictionary.classes) - {"10540832"}:
        raise ValueError("research rule coverage incomplete")
    audit, seen = [], set()
    for original in records:
        if not isinstance(original, dict) or {"source_fid", geom_col, *fields.values()} - set(original):
            raise ValueError("POI record missing required fields")
        fid = original["source_fid"]
        if type(fid) is not int or fid in seen:
            raise ValueError(f"invalid/duplicate source_fid in batch: {fid!r}")
        seen.add(fid)
        reasons, exclusions = [], []

        def reason(code, action, detail=""):
            reasons.append({"reason": code, "action": action, "detail": detail})
            if action == "exclude":
                exclusions.append(code)

        identity, raw_code = original[fields["ref_no"]], original[fields["raw_code"]]
        kind = _missing_kind(identity)
        valid_id = kind == "present" and type(identity) in (str, int, float)
        if not valid_id:
            reason("invalid_ref_no_" + (kind if kind != "present" else type(identity).__name__), "exclude")
        elif isinstance(identity, str) and identity != identity.strip():
            reason("ref_no_surrounding_whitespace_preserved", "review")
        key = _raw_key(identity) if type(identity) is not bool else None
        duplicate = key in context.duplicate_keys
        if fid in context.duplicate_rows and context.duplicate_rows[fid] != key:
            raise ValueError(f"source_fid/ID differs from full-source duplicate evidence: {fid}")
        if duplicate:
            if context.duplicate_rows.get(fid) != key:
                raise ValueError(f"duplicate ID row absent from complete conflict evidence: {fid}")
            valid_id = False
            reason("global_duplicate_ref_no", "exclude", str(context.duplicate_keys[key]))
        geo = _geometry_check(original[geom_col], schema["geometry"]["srs_id"])
        valid_geometry = geo["status"] == "valid_point"
        if not valid_geometry:
            reason("geometry_" + geo["status"], "exclude", geo["detail"])
        if geo["has_z"] or geo["has_m"]:
            reason("geometry_extra_dimensions_xy_only", "review")
        within = bool(context.coverage.covers(Point(geo["xy"]))) if valid_geometry else None
        if within is False:
            reason("outside_data_coverage", "exclude")
        elif within is None:
            reason("coverage_not_evaluated_invalid_geometry", "not_evaluated")
        facility = classify(raw_code, context.dictionary.classes, include_ev_charging=include_ev_charging)
        if not facility.include_in_count:
            reason("charging_facility_switch_excluded", "exclude")
        classification = research_classification(raw_code, dictionary=context.dictionary, rules=context.rules)
        accuracy = original[fields["accuracy"]]
        accuracy_code = str(accuracy)
        if accuracy_code in ("3", "4"):
            reason("low_position_accuracy_" + accuracy_code, "review", context.accuracy_labels[accuracy_code])
        elif accuracy_code not in context.accuracy_labels:
            reason("missing_or_unknown_position_accuracy", "review", json.dumps(_poi_typed(accuracy)))
        x_attr, y_attr = (_number(original[fields[k]]) for k in ("easting", "northing"))
        dx = dy = delta = None
        difference_status = "not_comparable"
        if valid_geometry and x_attr is not None and y_attr is not None:
            # Both coordinates belong to this verified BNG schema. Never use
            # deltas when a caller supplies different attribute CRS metadata.
            attribute_crs = original.get("attribute_crs", schema["crs"]["definition"])
            try:
                same_crs = CRS.from_user_input(attribute_crs).equals(crs)
            except (TypeError, ValueError, CRSError):
                same_crs = False
            if same_crs:
                dx, dy = geo["xy"][0] - x_attr, geo["xy"][1] - y_attr
                delta = math.hypot(dx, dy)
                if not all(math.isfinite(v) for v in (dx, dy, delta)):
                    dx = dy = delta = None
                    reason("coordinate_difference_overflow", "review")
                else:
                    difference_status = "comparable"
                    if delta != 0:
                        reason("geometry_attribute_difference", "review", repr(delta))
            else:
                reason("attribute_coordinate_crs_mismatch", "review")
        elif x_attr is None or y_attr is None:
            reason("invalid_attribute_coordinates", "review")
        else:
            reason("coordinate_difference_unavailable_geometry", "not_evaluated")
        base = valid_geometry and valid_id and within is True and facility.include_in_count
        included = base and classification["research_classification_include"]
        analysis_reasons = ["base:" + r for r in exclusions]
        if not classification["research_classification_include"]:
            analysis_reasons.append("classification:" + classification["classification_reason"])
        raw_fields = {name: _poi_typed(value) for name, value in original.items() if name != geom_col}
        blob = original[geom_col]
        result = {"source_path": context.source_path, "source_layer": context.layer, "source_fid": fid,
                  "raw_ref_no": _poi_typed(identity), "raw_classification": _poi_typed(raw_code),
                  "raw_fields": raw_fields, "raw_geometry": _poi_typed(blob),
                  "geometry_gpkg_hex": blob.hex() if isinstance(blob, bytes) else None,
                  "geometry": geo, "x": geo["xy"][0] if valid_geometry else None,
                  "y": geo["xy"][1] if valid_geometry else None, "crs": "EPSG:27700",
                  "geometry_valid": valid_geometry, "id_valid": valid_id, "global_duplicate_id": duplicate,
                  "within_data_coverage": within, "facility_switch_passed": facility.include_in_count,
                  "raw_position_accuracy": _poi_typed(accuracy), "position_accuracy_review": accuracy_code in ("3", "4"),
                  "coordinate_difference_status": difference_status, "dx": dx, "dy": dy, "delta_m": delta,
                  "quality_reasons": reasons, "base_exclusion_reasons": exclusions,
                  "base_eligible": bool(base), "analysis_include": bool(included), **classification,
                  "analysis_exclusion_reason": analysis_reasons,
                  "time_audit": {"supply_date_is_release_date": "supply_date" in original,
                                 "topo_toid_version_is_not_poi_lifecycle": "topo_toid_version" in original,
                                 "historical_existence_not_inferred": True,
                                 "lifecycle_status": "no_verified_lifecycle_field"}}
        audit.append(result)
    return {"base_pois": [r for r in audit if r["base_eligible"]],
            "analysis_pois": [r for r in audit if r["analysis_include"]], "audit": audit,
            "summary": poi_cleaning_summary(audit, include_ev_charging=include_ev_charging,
                                            rule_version=context.rules.version, group_dictionary=context.dictionary.groups)}


def poi_cleaning_summary(audit, *, include_ev_charging, rule_version, group_dictionary):
    n = len(audit)
    bases = sum(r["base_eligible"] for r in audit)
    includes = sum(r["analysis_include"] for r in audit)
    reviews = sum(r["classification_review"] for r in audit)
    anomaly_statuses = ("unknown_dictionary", "missing", "malformed")
    anomalies = sum(r["classification_status"] in anomaly_statuses for r in audit)
    native_counts = Counter(r["group_code"] for r in audit if r["analysis_include"])
    group_codes = sorted(set(group_dictionary) | set(native_counts))
    return {"schema_version": 2, "composition_mode": "native_group",
            "input_count": n, "base_eligible_count": bases, "base_excluded_count": n - bases,
            "analysis_included_count": includes, "analysis_excluded_count": n - includes,
            "base_eligible_research_excluded_count": bases - includes,
            "classification_review_count": reviews, "classification_review_proportion": reviews / n if n else None,
            "classification_anomaly_count": anomalies, "classification_anomaly_proportion": anomalies / n if n else None,
            "base_eligible_classification_review_count": sum(r["base_eligible"] and r["classification_review"] for r in audit),
            "base_exclusion_reasons": dict(Counter(x for r in audit for x in r["base_exclusion_reasons"])),
            "quality_reasons": dict(Counter(x["reason"] for r in audit for x in r["quality_reasons"])),
            "analysis_exclusion_reasons": dict(Counter(x for r in audit for x in r["analysis_exclusion_reason"])),
            "classification_status_counts": dict(Counter(r["classification_status"] for r in audit)),
            "research_decision_counts": dict(Counter(r["research_decision"] for r in audit)),
            "multiple_base_exclusion_reason_count": sum(len(r["base_exclusion_reasons"]) > 1 for r in audit),
            "native_group_counts": {g: native_counts[g] for g in group_codes},
            "native_group_names": {g: group_dictionary.get(g, {}).get("dictionary_name") for g in group_codes},
            "functional_group_pending_count": sum(r["analysis_include"] and r["functional_group_pending"] for r in audit),
            "functional_group_role": "audit_only", "unmatched_class_policy": "review",
            "include_ev_charging": include_ev_charging, "classification_rule_version": rule_version}


def poi_parquet_schema(rule_version):
    import pyarrow as pa
    fields = [("source_path", pa.string()), ("source_layer", pa.string()), ("source_fid", pa.int64()),
              ("ref_no_type", pa.string()), ("ref_no_value", pa.string()),
              ("raw_classification_type", pa.string()), ("raw_classification_value", pa.string()),
              ("raw_fields_json", pa.string()), ("geometry_gpkg", pa.binary()), ("geometry_wkb", pa.binary()),
              ("x", pa.float64()), ("y", pa.float64()), ("group_code", pa.string()),
              ("category_code", pa.string()), ("category_number", pa.string()), ("class_code", pa.string()),
              ("group_dictionary_name", pa.string()), ("category_dictionary_name", pa.string()),
              ("class_dictionary_name", pa.string()), ("classification_status", pa.string()),
              ("research_decision", pa.string()), ("classification_reason", pa.string()),
              ("base_eligible", pa.bool_()), ("analysis_include", pa.bool_()),
              ("poi_function_group", pa.string()), ("functional_group_pending", pa.bool_()),
              ("classification_review", pa.bool_()),
              ("analysis_exclusion_reason_json", pa.string()), ("classification_rule_version", pa.string())]
    return pa.schema(fields, metadata={b"crs": b"EPSG:27700", b"schema_version": b"2",
                                      b"composition_mode": b"native_group", b"functional_group_role": b"audit_only",
                                      b"classification_rule_version": rule_version.encode()})


def _poi_parquet_row(row):
    from shapely.geometry import Point
    result = {k: row[k] for k in ("source_path", "source_layer", "source_fid", "x", "y", "group_code", "category_code", "category_number", "class_code", "base_eligible", "analysis_include", "poi_function_group", "classification_review", "classification_rule_version")}
    result.update({k: row[k] for k in ("group_dictionary_name", "category_dictionary_name", "class_dictionary_name",
                                      "classification_status", "research_decision", "classification_reason", "functional_group_pending")})
    result.update(ref_no_type=row["raw_ref_no"]["type"], ref_no_value=row["raw_ref_no"]["value"],
                  raw_classification_type=row["raw_classification"]["type"], raw_classification_value=row["raw_classification"]["value"],
                  raw_fields_json=json.dumps(row["raw_fields"], ensure_ascii=False, allow_nan=False, sort_keys=True),
                  geometry_gpkg=bytes.fromhex(row["geometry_gpkg_hex"]), geometry_wkb=Point(row["x"], row["y"]).wkb,
                  analysis_exclusion_reason_json=json.dumps(row["analysis_exclusion_reason"], ensure_ascii=False))
    return result


class POICleaningError(ValueError):
    def __init__(self, report):
        self.report = report
        super().__init__("; ".join(report["errors"]))


def prepare_poi_sample(poi_path, *, data_root, code_root, output_dir, context_options,
                       limit=100, batch_size=100, include_ev_charging=False):
    """Bounded first-FID integration helper. It cannot start a full-source run.

    Generic clean_poi_batch remains reusable for future authorized streaming;
    this file driver deliberately accepts only limits <= 100 in step 11.
    """
    import pyarrow as pa
    import pyarrow.parquet as pq
    _positive_integer(limit, "limit")
    _positive_integer(batch_size, "batch_size")
    if limit > 100 or type(include_ev_charging) is not bool:
        raise ValueError("step11 sample driver requires limit <= 100 and a boolean switch; full-source cleaning is not authorized")
    root = project_path(".", root=data_root)
    output = _audit_path(output_dir, root)
    if output == root / "outputs":
        raise ValueError("output must be strictly inside outputs/")
    output.mkdir(parents=True, exist_ok=False)
    report = {"schema_version": 2, "composition_mode": "native_group", "status": "running", "validation_complete": False,
              "scope": "first_records_sample", "full_source_cleaning_run": False,
              "limit": limit, "batch_size": batch_size, "include_ev_charging": include_ev_charging,
              "output_dir": str(output), "errors": []}
    _json_write(output / "run_manifest.json", report)
    try:
        context = load_poi_cleaning_context(poi_path, data_root=root, code_root=code_root, **context_options)
        report.update(classification_rule_version=context.rules.version, upstream=context.provenance,
                      input_schema=context.schema, rules=context.rules.provenance)
        schema = poi_parquet_schema(context.rules.version)
        base_path, main_path, audit_path = (output / n for n in ("base_eligible_pois.parquet", "research_pois.parquet", "poi_audit.jsonl"))
        # Publication occurs only after writers close and independent read-back.
        temporary = [p.with_suffix(p.suffix + ".tmp") for p in (base_path, main_path, audit_path)]
        seen, all_audit = set(), []  # bounded by limit <= 100, never whole-source memory
        with ExitStack() as stack:
            base_writer = stack.enter_context(pq.ParquetWriter(temporary[0], schema))
            main_writer = stack.enter_context(pq.ParquetWriter(temporary[1], schema))
            stream = stack.enter_context(temporary[2].open("x", encoding="utf-8"))
            for batch in poi_batches(poi_path, data_root=root, layer=context.layer, fields=context.schema["field_mapping"], batch_size=batch_size, limit=limit):
                if len(seen) + len(batch) > limit:
                    raise ValueError("reader exceeded authorized sample limit")
                result = clean_poi_batch(batch, context=context, include_ev_charging=include_ev_charging)
                for row in result["audit"]:
                    if row["source_fid"] in seen:
                        raise ValueError("duplicate source_fid across batches")
                    seen.add(row["source_fid"])
                    all_audit.append(row)
                    stream.write(json.dumps(row, ensure_ascii=False, allow_nan=False, sort_keys=True) + "\n")
                for writer, key in ((base_writer, "base_pois"), (main_writer, "analysis_pois")):
                    writer.write_table(pa.Table.from_pylist([_poi_parquet_row(r) for r in result[key]], schema=schema))
        if len(all_audit) != min(limit, context.schema["source_count"]):
            raise ValueError("sample reader/audit count does not reconcile")
        restored_audit = [json.loads(line, object_pairs_hook=_unique_json_object) for line in temporary[2].read_text().splitlines()]
        # Tuples in geometry xy serialize as JSON lists; compare canonical JSON.
        canonical = json.loads(json.dumps(all_audit, ensure_ascii=False, allow_nan=False))
        if restored_audit != canonical:
            raise ValueError("complete POI audit read-back differs")
        for file, keep in ((temporary[0], "base_eligible"), (temporary[1], "analysis_include")):
            restored = pq.read_table(file)
            expected = [_poi_parquet_row(r) for r in all_audit if r[keep]]
            if not restored.schema.equals(schema, check_metadata=True) or restored.to_pylist() != expected:
                raise ValueError("POI Parquet schema/rows differ on read-back")
        summary = poi_cleaning_summary(all_audit, include_ev_charging=include_ev_charging,
                                       rule_version=context.rules.version, group_dictionary=context.dictionary.groups)
        if (sum(summary["native_group_counts"].values()) != summary["analysis_included_count"]
                or any(r["analysis_include"] and (not r["base_eligible"] or r["research_decision"] != "include"
                       or r["group_code"] not in context.dictionary.groups or r["class_code"] == "10540832") for r in all_audit)
                or any(not r["analysis_include"] and not r["analysis_exclusion_reason"] for r in all_audit)):
            raise ValueError("POI eligibility/native Group/reason invariants failed")
        # Reverify evidence/config/manifest/raw bytes after processing, before publication.
        after = load_poi_cleaning_context(poi_path, data_root=root, code_root=code_root, **context_options)
        if after.provenance != context.provenance or after.rules.coverage != context.rules.coverage:
            raise ValueError("inputs/upstream/rules changed during POI sample cleaning")
        for key, fingerprint_key, name in (("config_path", "config_fingerprint", "research_rules_snapshot.yaml"),
                                            ("rules_file", "rules_fingerprint", "research_rules_snapshot.csv")):
            source = project_path(context.rules.provenance[key], root=code_root)
            expected = context.rules.provenance[fingerprint_key]
            snapshot = output / (name + ".tmp")
            with _open_regular(source) as stream, snapshot.open("xb") as target:
                target.write(stream.read())
            if _artifact_fingerprint(snapshot) != expected or _artifact_fingerprint(source) != expected:
                raise ValueError("research rule snapshot/source differs from verified rules")
            snapshot.replace(output / name)
        for source, target in zip(temporary, (base_path, main_path, audit_path)):
            source.replace(target)
        _json_write(output / "poi_summary.json", summary)
        _json_write(output / "rule_coverage.json", context.rules.coverage)
        validation = {"status": "passed", "scope": "first_records_sample", "full_source_run": False,
                      "counts_conserved": True, "audit_one_per_source_record": True,
                      "parquet_schema_and_rows_verified": True, "native_groups_sum_verified": True,
                      "composition_mode": "native_group", "functional_group_role": "audit_only",
                      "all_excluded_rows_have_reasons": True, "input_bytes_unchanged": True,
                      "upstream_and_rules_unchanged": True, "summary": summary}
        _json_write(output / "validation_report.json", validation)
        report.update(status="completed", validation_complete=True, input_bytes_unchanged=True, summary=summary,
                      artifacts={p.name: _artifact_fingerprint(p) for p in sorted(output.iterdir()) if p.name != "run_manifest.json"})
    except Exception as error:
        report.update(status="failed", validation_complete=False)
        report["errors"].append(f"{type(error).__name__}: {error}")
        _json_write(output / "run_manifest.json", report)
        raise POICleaningError(report) from error
    _json_write(output / "run_manifest.json", report)
    return report
